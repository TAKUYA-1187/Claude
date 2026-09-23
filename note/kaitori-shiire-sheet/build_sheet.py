"""note で配布する「買取せどり 仕入れ判定・入金管理シート」を生成する。

    python note/kaitori-shiire-sheet/build_sheet.py

出力: note/kaitori-shiire-sheet/kaitori_shiire_sheet.xlsx
判定基準の初期値は運営者の実際の設定（利益500円以上・利益率15%以上・送料600円/箱）。
"""
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).with_name("kaitori_shiire_sheet.xlsx")
FONT = "Yu Gothic"
FIRST, LAST = 6, 505  # 仕入れ判定シートのデータ行（500件分）
STATUSES = ["未発送", "発送済", "査定済", "入金済", "見送り"]

BLUE = Font(name=FONT, color="0000FF")
BLACK = Font(name=FONT)
BOLD = Font(name=FONT, bold=True)
TITLE = Font(name=FONT, bold=True, size=14)
HEAD = Font(name=FONT, bold=True, color="FFFFFF")
HEAD_IN = PatternFill("solid", fgColor="0F6B52")   # 入力列の見出し
HEAD_OUT = PatternFill("solid", fgColor="5B636E")  # 自動計算列の見出し
INPUT = PatternFill("solid", fgColor="FFF9DB")     # 入力セル
OK_FILL = PatternFill("solid", fgColor="E6F2EE")
WEAK_FILL = PatternFill("solid", fgColor="FFF4CC")
NG_FILL = PatternFill("solid", fgColor="FBEEE6")
THIN = Side(style="thin", color="D0CCC4")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
YEN = '#,##0;[Red]-#,##0;"-"'
PCT = '0.0%;[Red]-0.0%;"-"'
DATE = "yyyy/mm/dd"
MONTH = "yyyy年m月"

S = "'設定'!"
MIN_PROFIT, MIN_RATE, SHIP, DROP = f"{S}$C$4", f"{S}$C$5", f"{S}$C$6", f"{S}$C$7"
J = "'仕入れ判定'!"


def col_range(col: str) -> str:
    return f"{J}${col}${FIRST}:${col}${LAST}"


def build_settings(wb: Workbook) -> None:
    ws = wb.create_sheet("設定")
    ws["A1"] = "判定基準の設定"
    ws["A1"].font = TITLE
    ws["A2"] = "黄色のセルだけ書き換えてください。変えると全シートの判定が変わります。"
    ws["A2"].font = BLACK
    rows = [
        ("最低利益（円）", 500, YEN, "1件あたり、これ未満の利益なら見送る。初期値は運営者の実際の基準。"),
        ("最低利益率", 0.15, PCT, "利益 ÷ 仕入れ価格。これ未満なら見送る。初期値は運営者の実際の基準。"),
        ("送料／1箱（円）", 600, YEN, "箱ごとに違う場合は、仕入れ判定シートの「箱の送料」列に入れると優先される。"),
        ("買取価格の下落シナリオ", 0.10, PCT, "発送から査定までに買取価格がこの割合だけ下がったときの利益を併記する。"),
    ]
    for i, (label, value, fmt, note) in enumerate(rows, start=4):
        ws.cell(i, 2, label).font = BOLD
        c = ws.cell(i, 3, value)
        c.font, c.fill, c.number_format, c.border = BLUE, INPUT, fmt, BOX
        ws.cell(i, 4, note).font = BLACK
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 80


def build_judge(wb: Workbook) -> None:
    ws = wb.create_sheet("仕入れ判定")
    ws["A1"] = "仕入れ判定・入金管理"
    ws["A1"].font = TITLE
    ws["A2"] = (
        "緑の見出し＝入力する列／灰色の見出し＝自動計算。同じ箱で送る商品には同じ「箱ID」を付けると、"
        "送料が点数で自動的に割られます。"
    )
    ws["A3"] = (
        "【例】の3行は見本です（同梱で送料が割られる例と、利益率不足で見送る例）。"
        "使い始めるときは行ごと消すか、上書きしてください。"
    )
    for c in ("A2", "A3"):
        ws[c].font = BLACK

    # (見出し, 入力列か, 幅, 書式, 数式テンプレート)
    cols = [
        ("仕入れ日", True, 11, DATE, None),
        ("商品名", True, 30, None, None),
        ("JAN", True, 15, "@", None),
        ("仕入れ先", True, 11, None, None),
        ("買取店", True, 12, None, None),
        ("箱ID", True, 8, None, None),
        ("仕入れ価格\n（支払額）", True, 11, YEN, None),
        ("ポイント\n還元率", True, 9, PCT, None),
        ("買取価格\n（提示額）", True, 11, YEN, None),
        ("その他\nコスト", True, 9, YEN, None),
        ("箱の送料\n（空欄=設定値）", True, 12, YEN, None),
        ("同梱点数", False, 8, "0", '=IF($G{r}="","",IF($F{r}="",1,MAX(1,COUNTIF($F${a}:$F${b},$F{r})-COUNTIFS($F${a}:$F${b},$F{r},$S${a}:$S${b},"見送り"))))'),
        ("1点あたり\n送料", False, 10, YEN, f'=IF($G{{r}}="","",IF($K{{r}}="",{SHIP},$K{{r}})/$L{{r}})'),
        ("見込み利益", False, 10, YEN, '=IF(OR($G{r}="",$I{r}=""),"",$I{r}-$G{r}-$M{r}-$J{r})'),
        ("利益率", False, 8, PCT, '=IF($N{r}="","",IF($G{r}>0,$N{r}/$G{r},0))'),
        ("ポイント\n（別枠）", False, 9, YEN, '=IF($G{r}="","",$G{r}*$H{r})'),
        ("下落時\n利益", False, 9, YEN, f'=IF($N{{r}}="","",$N{{r}}-$I{{r}}*{DROP})'),
        (
            "判定",
            False,
            14,
            None,
            f'=IF($N{{r}}="","",IF(AND($N{{r}}>={MIN_PROFIT},$O{{r}}>={MIN_RATE}),'
            f'IF($Q{{r}}>=0,"仕入れOK","OK・下落に弱い"),"見送り"))',
        ),
        ("状態", True, 9, None, None),
        ("発送日", True, 11, DATE, None),
        ("確定\n買取額", True, 10, YEN, None),
        ("入金日", True, 11, DATE, None),
        ("確定利益", False, 10, YEN, '=IF(OR($G{r}="",$U{r}=""),"",$U{r}-$G{r}-$M{r}-$J{r})'),
        ("減額", False, 9, YEN, '=IF(OR($I{r}="",$U{r}=""),"",$I{r}-$U{r})'),
        ("発送→入金\n日数", False, 9, "0", '=IF(OR($T{r}="",$V{r}=""),"",$V{r}-$T{r})'),
        # 状態が入った行＝実際に仕入れた行。空欄・見送りは判定だけした候補として集計しない
        ("仕入れ済み", False, 8, "0", '=IF(OR($S{r}="未発送",$S{r}="発送済",$S{r}="査定済",$S{r}="入金済"),1,0)'),
        ("確定済み", False, 8, "0", '=IF($U{r}="",0,1)'),
    ]
    header_row = FIRST - 1
    for idx, (name, is_input, width, fmt, formula) in enumerate(cols, start=1):
        letter = get_column_letter(idx)
        h = ws.cell(header_row, idx, name)
        h.font, h.fill, h.border = HEAD, HEAD_IN if is_input else HEAD_OUT, BOX
        h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[letter].width = width
        for r in range(FIRST, LAST + 1):
            c = ws.cell(r, idx)
            c.border = BOX
            if fmt:
                c.number_format = fmt
            if formula:
                c.value = formula.format(r=r, a=FIRST, b=LAST)
                c.font = BLACK
            else:
                c.font = BLUE
                c.fill = INPUT
    ws.row_dimensions[header_row].height = 36

    ws["N5"].comment = Comment("買取価格 − 仕入れ価格 − 1点あたり送料 − その他コスト。ポイントは含めない。", "sheet")
    ws["Q5"].comment = Comment("設定シートの下落シナリオ（初期値10%）だけ買取価格が下がった場合の利益。", "sheet")
    ws["R5"].comment = Comment(
        "見込み利益と利益率が設定の基準以上なら OK。さらに下落時利益がマイナスなら「下落に弱い」。", "sheet"
    )

    # 見本：箱A1に3点の候補。イヤホンは利益率不足で見送り→箱は2点になり、送料600円が300円ずつに割られる。
    # 周辺機器は単品で送ると利益200円で見送りだが、同梱なら基準を満たす。
    examples = [
        (date(2026, 10, 1), "【例】ゲームソフトA", "4900000000001", "楽天市場", "買取店X", "A1", 3800, 0.05, 5000, 0, "未発送"),
        (date(2026, 10, 1), "【例】イヤホンB", "4900000000002", "Yahoo!", "買取店X", "A1", 9200, 0.10, 10500, 0, "見送り"),
        (date(2026, 10, 2), "【例】周辺機器C", "4900000000003", "Amazon", "買取店X", "A1", 2400, 0.01, 3200, 0, "未発送"),
    ]
    for i, (*row, status) in enumerate(examples):
        r = FIRST + i
        for j, v in enumerate(row, start=1):
            ws.cell(r, j, v)
        ws.cell(r, 19, status)

    dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True)
    dv.error, dv.errorTitle = "リストから選んでください", "状態"
    ws.add_data_validation(dv)
    dv.add(f"S{FIRST}:S{LAST}")

    rng = f"R{FIRST}:R{LAST}"
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'$R{FIRST}="仕入れOK"'], fill=OK_FILL))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'$R{FIRST}="OK・下落に弱い"'], fill=WEAK_FILL))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'$R{FIRST}="見送り"'], fill=NG_FILL))
    ws.freeze_panes = ws.cell(FIRST, 3)
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(cols))}{LAST}"


def build_summary(wb: Workbook) -> None:
    ws = wb.create_sheet("集計")
    ws["A1"] = "お金の流れと月別の成績"
    ws["A1"].font = TITLE
    st, cost, buy, fixed = col_range("S"), col_range("G"), col_range("I"), col_range("U")

    ws["A3"] = "いま立て替えている現金"
    ws["A3"].font = BOLD
    lines = [
        ("仕入れ済みで未入金の支払額", "+".join(f'SUMIFS({cost},{st},"{s}")' for s in ("未発送", "発送済", "査定済"))),
        ("入金待ちの見込み額（発送済は提示額、査定済は確定額）", f'SUMIFS({buy},{st},"発送済")+SUMIFS({fixed},{st},"査定済")'),
        ("未発送の点数", f'COUNTIFS({st},"未発送")'),
    ]
    for i, (label, f) in enumerate(lines, start=4):
        ws.cell(i, 1, label).font = BLACK
        c = ws.cell(i, 4, "=" + f)
        c.font, c.border = BLACK, BOX
        c.number_format = "0" if "点数" in label else YEN
    ws["E4"] = "カードの支払日までに、この金額が戻ってくるかを確認する。"
    ws["E4"].font = BLACK

    ws["A9"] = "月別の成績（A列の月は書き換えてよい）"
    ws["A9"].font = BOLD
    heads = ["月", "仕入れ件数", "仕入れ額", "見込み利益", "確定利益", "減額合計", "確定利益率"]
    for j, h in enumerate(heads, start=1):
        c = ws.cell(10, j, h)
        c.font, c.fill, c.border = HEAD, HEAD_OUT, BOX
        c.alignment = Alignment(horizontal="center")
    day, prof, fprof, cut = col_range("A"), col_range("N"), col_range("W"), col_range("X")
    bought, done = col_range("Z"), col_range("AA")
    for k in range(12):
        r = 11 + k
        m = ws.cell(r, 1, date(2026, 10, 1) if k == 0 else f"=EDATE(A{r - 1},1)")
        m.number_format, m.border = MONTH, BOX
        m.font = BLUE if k == 0 else BLACK
        if k == 0:
            m.fill = INPUT
        span = f'{day},">="&$A{r},{day},"<"&EDATE($A{r},1)'
        formulas = [
            f"=SUMIFS({bought},{span})",
            f"=SUMIFS({cost},{span},{bought},1)",
            f"=SUMIFS({prof},{span},{bought},1)",
            f"=SUMIFS({fprof},{span},{done},1)",
            f"=SUMIFS({cut},{span},{done},1)",
            f"=IFERROR(E{r}/SUMIFS({cost},{span},{done},1),0)",
        ]
        fmts = ["0", YEN, YEN, YEN, YEN, PCT]
        for j, (f, fmt) in enumerate(zip(formulas, fmts), start=2):
            c = ws.cell(r, j, f)
            c.font, c.number_format, c.border = BLACK, fmt, BOX
    ws["A24"] = "確定利益率＝確定利益 ÷ 確定買取額が入った商品の仕入れ額。見込み利益との差が、減額・送料の見込み違いの大きさ。"
    ws["A24"].font = BLACK

    ws["A26"] = "買取店別の減額（B列に買取店名を入れる。仕入れ判定シートの「買取店」と同じ表記で）"
    ws["A26"].font = BOLD
    heads = ["", "買取店", "査定済みの件数", "提示額の合計", "減額合計", "減額率"]
    for j, h in enumerate(heads, start=1):
        if h:
            c = ws.cell(27, j, h)
            c.font, c.fill, c.border = HEAD, HEAD_OUT, BOX
    shop = col_range("E")
    for k in range(8):
        r = 28 + k
        name = ws.cell(r, 2, "買取店X" if k == 0 else None)
        name.font, name.fill, name.border = BLUE, INPUT, BOX
        cells = [
            (f'=IF($B{r}="","",SUMIFS({done},{shop},$B{r}))', "0"),
            (f'=IF($B{r}="","",SUMIFS({buy},{shop},$B{r},{done},1))', YEN),
            (f'=IF($B{r}="","",SUMIFS({cut},{shop},$B{r},{done},1))', YEN),
            (f'=IF(OR($B{r}="",N(D{r})=0),"",E{r}/D{r})', PCT),
        ]
        for j, (f, fmt) in enumerate(cells, start=3):
            c = ws.cell(r, j, f)
            c.font, c.number_format, c.border = BLACK, fmt, BOX
    ws["A37"] = "減額率が高い買取店は、梱包の見直しか、送り先の変更を検討する。"
    ws["A37"].font = BLACK
    ws.column_dimensions["A"].width = 14
    for letter in "BCDEFG":
        ws.column_dimensions[letter].width = 15


def build_guide(wb: Workbook) -> None:
    ws = wb.active
    ws.title = "使い方"
    lines = [
        ("買取せどり 仕入れ判定・入金管理シート", TITLE),
        ("", BLACK),
        ("入力するのは黄色のセル（青い文字）だけです。灰色の見出しの列と、集計シートは自動で計算されます。", BLACK),
        ("", BLACK),
        ("1. 設定シート：最低利益・最低利益率・送料／1箱・下落シナリオを、自分の基準に合わせる", BOLD),
        ("2. 仕入れ判定シート：仕入れる前に、仕入れ価格と買取価格を入れて「判定」を見る", BOLD),
        ("   ・同じ箱で送る商品には同じ箱ID（例：A1）を付ける。送料が点数で自動的に割られる（状態が「見送り」の行は点数に数えない）", BLACK),
        ("   ・「OK・下落に弱い」は、買取価格が下落シナリオ分下がると赤字になる商品。点数を控えるか、すぐ発送する", BLACK),
        ("   ・ポイントは判定に入れない（付与の時期・上限があるため）。別枠の列で確認する", BLACK),
        ("3. 仕入れたら「状態」を 未発送 → 発送済 → 査定済 → 入金済 と更新し、確定買取額と日付を入れる", BOLD),
        ("   ・状態が空欄の行は「判定しただけの候補」として、集計に数えない。見送った商品は「見送り」にしておくと後で見返せる", BLACK),
        ("4. 集計シート：立て替えている現金、月別の見込み利益と確定利益、買取店別の減額率を確認する", BOLD),
        ("", BLACK),
        ("判定のルール", BOLD),
        ("仕入れOK：見込み利益 ≥ 最低利益、かつ 利益率 ≥ 最低利益率、かつ 下落時利益 ≥ 0", BLACK),
        ("OK・下落に弱い：基準は満たすが、下落時利益がマイナス", BLACK),
        ("見送り：見込み利益か利益率が基準未満", BLACK),
        ("", BLACK),
        ("注意", BOLD),
        ("・仕入れ価格は、クーポン・送料を反映した支払額を入れる", BLACK),
        ("・買取価格は日々変わります。判定は入力した時点の数字での計算で、利益を保証するものではありません", BLACK),
        ("・初期値の基準（利益500円・利益率15%・送料600円/箱）は運営者の実際の設定です。自分の送料・資金に合わせて変えてください", BLACK),
    ]
    for i, (text, font) in enumerate(lines, start=1):
        c = ws.cell(i, 1, text)
        c.font = font
    ws.column_dimensions["A"].width = 120


def main() -> None:
    wb = Workbook()
    build_guide(wb)
    build_settings(wb)
    build_judge(wb)
    build_summary(wb)
    wb.calculation.fullCalcOnLoad = True  # 開いたときに必ず再計算させる
    wb.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
