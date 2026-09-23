"""note 販売の数字を記録する表を作る。

    python note/kpi/build_tracker.py

出力: note/kpi/note_kpi_tracker.xlsx
毎週オーナーから届く数字（有料記事ごとのビュー・購入・高評価・返金、メンバー数）を入れ、
月100万円（手取り）までの進み具合と、商品ごとの購入率を見る。
"""
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).with_name("note_kpi_tracker.xlsx")
FONT = "Yu Gothic"
BLUE, BLACK = Font(name=FONT, color="0000FF"), Font(name=FONT)
BOLD, TITLE = Font(name=FONT, bold=True), Font(name=FONT, bold=True, size=14)
HEAD = Font(name=FONT, bold=True, color="FFFFFF")
HEAD_IN, HEAD_OUT = PatternFill("solid", fgColor="0F6B52"), PatternFill("solid", fgColor="5B636E")
INPUT = PatternFill("solid", fgColor="FFF9DB")
THIN = Side(style="thin", color="D0CCC4")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
YEN, PCT, DATE, MONTH = '#,##0;[Red]-#,##0;"-"', '0.0%;[Red]-0.0%;"-"', "yyyy/mm/dd", "yyyy年m月"
W_FIRST, W_LAST = 5, 504  # 週次シートのデータ行
RATE, GOAL = "'設定'!$C$4", "'設定'!$C$5"

PRODUCTS = [
    ("P1", "買取せどりの仕入れ判定・入金管理シート", 980),
    ("P2", "Amazon物販 利益・現金・在庫の監査キット", 2980),
    ("P3", "中国OEM 見積比較・検品・補償交渉キット", 4980),
    ("P4", "新商品の GO／NO-GO 判定シート", 2980),
    ("P5", "物販の利益管理 完全パック", 14800),
]


def header(ws, row, names, inputs):
    for j, (name, is_in) in enumerate(zip(names, inputs), start=1):
        c = ws.cell(row, j, name)
        c.font, c.fill, c.border = HEAD, HEAD_IN if is_in else HEAD_OUT, BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 32


def cell(ws, r, c, value, fmt=None, is_input=False):
    x = ws.cell(r, c, value)
    x.border = BOX
    x.font = BLUE if is_input else BLACK
    if is_input:
        x.fill = INPUT
    if fmt:
        x.number_format = fmt
    return x


def build(wb: Workbook) -> None:
    guide = wb.active
    guide.title = "使い方"
    for i, (text, font) in enumerate(
        [
            ("note 販売 KPI 記録表", TITLE),
            ("", BLACK),
            ("黄色のセル（青い文字）だけ入力します。ほかは自動で計算されます。", BLACK),
            ("1. 商品シート：note に公開したら、公開日と価格を入れる（値上げしたら価格を書き換える）", BOLD),
            ("2. 週次シート：毎週、商品ごとに1行。note のダッシュボードの「記事ごとのビュー」と、売上履歴の購入数を入れる", BOLD),
            ("3. 月次シート：月末に、フォロワー数・全体ビュー・メンバー数を入れる。単品の売上は週次から自動で集計される", BOLD),
            ("", BLACK),
            ("見方", BOLD),
            ("・購入率＝購入数 ÷ 有料記事のビュー。1%未満なら無料部分（何が手に入るか）を直す。2%以上なら値上げを検討する", BLACK),
            ("・高評価率＝高評価 ÷ 購入数。40%未満なら、無料部分の説明と中身のズレを疑う", BLACK),
            ("・手取りは設定シートの手取り率（初期値：カード決済の85.5%）で見積もる。実際の入金額は note の売上管理で確認する", BLACK),
        ],
        start=1,
    ):
        guide.cell(i, 1, text).font = font
    guide.column_dimensions["A"].width = 110

    st = wb.create_sheet("設定")
    st["A1"], st["A1"].font = "設定", TITLE
    for r, (label, value, fmt, note) in enumerate(
        [
            ("手取り率", 0.855, PCT, "クレジットカード決済：事務手数料5%＋残額にプラットフォーム利用料10%。キャリア決済が多いと下がる"),
            ("目標（手取り・月）", 1_000_000, YEN, "note/README.md の3章"),
        ],
        start=4,
    ):
        st.cell(r, 2, label).font = BOLD
        cell(st, r, 3, value, fmt, True)
        st.cell(r, 4, note).font = BLACK
    st.column_dimensions["B"].width, st.column_dimensions["C"].width, st.column_dimensions["D"].width = 20, 14, 90

    pr = wb.create_sheet("商品")
    pr["A1"], pr["A1"].font = "商品一覧", TITLE
    header(pr, 3, ["ID", "商品名", "価格（円）", "公開日", "note の URL"], [True] * 5)
    for i, (pid, name, price) in enumerate(PRODUCTS, start=4):
        for j, (v, fmt) in enumerate([(pid, None), (name, None), (price, YEN), (None, DATE), (None, None)], start=1):
            cell(pr, i, j, v, fmt, True)
    for r in range(4 + len(PRODUCTS), 4 + 15):
        for j, fmt in enumerate([None, None, YEN, DATE, None], start=1):
            cell(pr, r, j, None, fmt, True)
    for letter, w in zip("ABCDE", (6, 44, 11, 12, 50)):
        pr.column_dimensions[letter].width = w

    wk = wb.create_sheet("週次")
    wk["A1"], wk["A1"].font = "週次の記録（商品ごとに1行）", TITLE
    wk["A2"] = "記入例の1行は消してかまいません。"
    names = ["週の開始日", "商品ID", "有料記事\nビュー", "購入数", "高評価", "返金", "価格", "売上", "手取り見込み", "購入率", "高評価率"]
    header(wk, W_FIRST - 1, names, [True] * 6 + [False] * 5)
    ids = "'商品'!$A$4:$A$18"
    prices = "'商品'!$C$4:$C$18"
    for r in range(W_FIRST, W_LAST + 1):
        for j, fmt in enumerate([DATE, None, "#,##0", "0", "0", "0"], start=1):
            cell(wk, r, j, None, fmt, True)
        cell(wk, r, 7, f'=IF($B{r}="","",IFERROR(INDEX({prices},MATCH($B{r},{ids},0)),""))', YEN)
        cell(wk, r, 8, f'=IF(OR($B{r}="",$G{r}=""),"",$G{r}*(N($D{r})-N($F{r})))', YEN)
        cell(wk, r, 9, f'=IF($H{r}="","",$H{r}*{RATE})', YEN)
        cell(wk, r, 10, f'=IF(N($C{r})=0,"",N($D{r})/$C{r})', PCT)
        cell(wk, r, 11, f'=IF(N($D{r})=0,"",N($E{r})/$D{r})', PCT)
    for j, v in enumerate([date(2026, 9, 28), "P1", 300, 3, 1, 0], start=1):
        wk.cell(W_FIRST, j, v)
    dv = DataValidation(type="list", formula1=ids, allow_blank=True)
    wk.add_data_validation(dv)
    dv.add(f"B{W_FIRST}:B{W_LAST}")
    for letter, w in zip("ABCDEFGHIJK", (12, 8, 10, 8, 8, 8, 10, 11, 12, 9, 9)):
        wk.column_dimensions[letter].width = w
    wk.freeze_panes = f"A{W_FIRST}"

    mo = wb.create_sheet("月次")
    mo["A1"], mo["A1"].font = "月次の成績と、目標までの距離", TITLE
    names = ["月", "フォロワー", "全体ビュー", "メンバー数", "メンバー\n月額", "単品の売上", "メンバーの売上", "売上合計", "手取り見込み", "目標の達成率", "単品の購入数", "単品の購入率"]
    header(mo, 3, names, [True, True, True, True, True] + [False] * 7)
    day, sales, views, buys = (f"'週次'!${c}${W_FIRST}:${c}${W_LAST}" for c in "AHCD")
    for k in range(18):
        r = 4 + k
        cell(mo, r, 1, date(2026, 9, 1) if k == 0 else f"=EDATE(A{r - 1},1)", MONTH, k == 0)
        for j, fmt in ((2, "#,##0"), (3, "#,##0"), (4, "#,##0"), (5, YEN)):
            cell(mo, r, j, 1480 if j == 5 else None, fmt, True)
        span = f'{day},">="&$A{r},{day},"<"&EDATE($A{r},1)'
        cell(mo, r, 6, f"=SUMIFS({sales},{span})", YEN)
        cell(mo, r, 7, f"=N(D{r})*N(E{r})", YEN)
        cell(mo, r, 8, f"=F{r}+G{r}", YEN)
        cell(mo, r, 9, f"=H{r}*{RATE}", YEN)
        cell(mo, r, 10, f"=IF({GOAL}>0,I{r}/{GOAL},0)", PCT)
        cell(mo, r, 11, f"=SUMIFS({buys},{span})", "#,##0")
        cell(mo, r, 12, f'=IFERROR(K{r}/SUMIFS({views},{span}),"")', PCT)
    mo.cell(23, 1, "週の開始日がその月に入る週を、その月の売上として数える。メンバーの売上は、メンバー数×月額の概算。").font = BLACK
    for letter, w in zip("ABCDEFGHIJKL", (11, 10, 10, 9, 9, 11, 12, 11, 12, 10, 10, 10)):
        mo.column_dimensions[letter].width = w
    mo.freeze_panes = "B4"


def main() -> None:
    wb = Workbook()
    build(wb)
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
