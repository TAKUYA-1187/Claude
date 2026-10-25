"""note で配布する「Amazon物販 利益・現金・在庫の監査キット」を生成する。

    python note/amazon-audit-kit/build_kit.py

出力: note/amazon-audit-kit/amazon_audit_kit.xlsx
月ごとの数字を入れると、売上から利益までの間でお金が消えた場所（原価・Amazon手数料・返品・広告）と、
1個あたりの利益、在庫の月数、現金の増減を出す。見本の数字は架空。
"""
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

OUT = Path(__file__).with_name("amazon_audit_kit.xlsx")
FONT = "Yu Gothic"
BLUE, BLACK = Font(name=FONT, color="0000FF"), Font(name=FONT)
BOLD, TITLE = Font(name=FONT, bold=True), Font(name=FONT, bold=True, size=14)
HEAD = Font(name=FONT, bold=True, color="FFFFFF")
HEAD_IN, HEAD_OUT = PatternFill("solid", fgColor="0F6B52"), PatternFill("solid", fgColor="5B636E")
INPUT = PatternFill("solid", fgColor="FFF9DB")
RED_FILL, YEL_FILL, OK_FILL = (PatternFill("solid", fgColor=c) for c in ("FBE3DC", "FFF4CC", "E6F2EE"))
THIN = Side(style="thin", color="D0CCC4")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
YEN, PCT, NUM, MONTHS = '#,##0;[Red]-#,##0;"-"', '0.0%;[Red]-0.0%;"-"', '#,##0;[Red]-#,##0;"-"', '0.0;[Red]-0.0;"-"'
MONTH = "yyyy年m月"
FIRST, LAST = 6, 29  # 月次シートのデータ行（24か月分）
S = "'設定'!"
TH_MARGIN, TH_RETURN, TH_STOCK = f"{S}$C$4", f"{S}$C$5", f"{S}$C$6"

# (見出し, 入力か, 幅, 書式, 数式)。数式中の {r} は行番号
COLS = [
    ("月", True, 10, MONTH, None),                                   # A
    ("売上\n（税込）", True, 12, YEN, None),                         # B
    ("販売個数", True, 9, NUM, None),                                 # C
    ("販売\n手数料", True, 10, YEN, None),                           # D
    ("FBA配送\n手数料", True, 10, YEN, None),                        # E
    ("在庫\n保管手数料", True, 10, YEN, None),                       # F
    ("その他の\nAmazon費用", True, 10, YEN, None),                   # G
    ("広告費", True, 10, YEN, None),                                  # H
    ("返金額", True, 10, YEN, None),                                  # I
    ("返品\n個数", True, 7, NUM, None),                               # J
    ("うち\n再販不可", True, 7, NUM, None),                           # K
    ("1個あたり\n着地原価", True, 10, YEN, None),                    # L
    ("月末の\n在庫個数", True, 9, NUM, None),                         # M
    ("仕入れの\n支払額", True, 11, YEN, None),                        # N
    ("Amazonから\nの入金額", True, 11, YEN, None),                    # O
    ("Amazon以外\nの支払い", True, 10, YEN, None),                   # P
    ("実質売上", False, 12, YEN, '=IF($B{r}="","",$B{r}-$I{r})'),                                          # Q
    ("原価", False, 11, YEN, '=IF($B{r}="","",($C{r}-($J{r}-$K{r}))*$L{r})'),                              # R
    ("Amazon\n費用合計", False, 11, YEN, '=IF($B{r}="","",$D{r}+$E{r}+$F{r}+$G{r})'),                     # S
    ("広告前\n利益", False, 11, YEN, '=IF($B{r}="","",$Q{r}-$R{r}-$S{r})'),                                # T
    ("利益", False, 11, YEN, '=IF($B{r}="","",$T{r}-$H{r})'),                                               # U
    ("1個あたり\n利益", False, 10, YEN, '=IF($B{r}="","",IF($C{r}-$J{r}>0,$U{r}/($C{r}-$J{r}),""))'),      # V
    ("広告前\n利益率", False, 9, PCT, '=IF(N($Q{r})=0,"",$T{r}/$Q{r})'),                                   # W
    ("原価率", False, 8, PCT, '=IF(N($Q{r})=0,"",$R{r}/$Q{r})'),                                            # X
    ("Amazon\n費用率", False, 8, PCT, '=IF(N($B{r})=0,"",$S{r}/$B{r})'),                                   # Y
    ("広告\n比率", False, 8, PCT, '=IF(N($B{r})=0,"",$H{r}/$B{r})'),                                        # Z
    ("返品率", False, 8, PCT, '=IF(N($C{r})=0,"",$J{r}/$C{r})'),                                            # AA
    ("在庫\n月数", False, 7, MONTHS, '=IF($B{r}="","",IF($C{r}-$J{r}>0,$M{r}/($C{r}-$J{r}),""))'),          # AB
    ("現金の\n増減", False, 11, YEN, '=IF($B{r}="","",$O{r}-$N{r}-$P{r})'),                                # AC
    (
        "判定",
        False,
        12,
        None,
        f'=IF($B{{r}}="","",IF(N($U{{r}})<0,"赤字",IF(OR(N($W{{r}})<{TH_MARGIN},N($AA{{r}})>{TH_RETURN},N($AB{{r}})>{TH_STOCK}),"注意","OK")))',
    ),                                                                                                      # AD
]

# 見本（架空）：原価率が高く、広告前の利益が薄い商品。秋に在庫を積みすぎて現金が減る
SAMPLE = [
    # 月, 売上, 個数, 販売手数料, FBA, 保管, その他, 広告, 返金, 返品, 再販不可, 着地原価, 月末在庫, 仕入れ支払, 入金, 他支払
    (date(2026, 4, 1), 1_380_000, 100, 207_000, 132_000, 9_000, 4_000, 52_000, 41_400, 3, 1, 8_200, 420, 0, 980_000, 52_000),
    (date(2026, 5, 1), 1_518_000, 110, 227_700, 145_200, 10_000, 4_000, 60_000, 55_200, 4, 2, 8_200, 310, 0, 1_090_000, 60_000),
    (date(2026, 6, 1), 1_656_000, 120, 248_400, 158_400, 9_500, 4_000, 66_000, 41_400, 3, 1, 8_200, 190, 2_460_000, 1_190_000, 66_000),
    (date(2026, 7, 1), 1_794_000, 130, 269_100, 171_600, 12_000, 4_000, 78_000, 69_000, 5, 3, 8_400, 560, 0, 1_260_000, 78_000),
    (date(2026, 8, 1), 1_656_000, 120, 248_400, 158_400, 15_000, 4_000, 84_000, 82_800, 6, 3, 8_400, 440, 0, 1_170_000, 84_000),
    (date(2026, 9, 1), 1_380_000, 100, 207_000, 132_000, 18_000, 4_000, 90_000, 55_200, 4, 2, 8_400, 900, 4_200_000, 990_000, 90_000),
]


def head(ws, row, col, text, is_input):
    c = ws.cell(row, col, text)
    c.font, c.fill, c.border = HEAD, HEAD_IN if is_input else HEAD_OUT, BOX
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def build_guide(wb):
    ws = wb.active
    ws.title = "使い方"
    lines = [
        ("Amazon物販 利益・現金・在庫の監査キット", TITLE),
        ("", BLACK),
        ("入力するのは黄色のセル（青い文字）だけです。ほかは自動で計算されます。【見本】の6か月分は架空の数字です。使うときは上書きしてください。", BLACK),
        ("", BLACK),
        ("1. 月次シート：1か月1行。数字はセラーセントラルの「ペイメント」→「日付範囲別レポート」と、自分の仕入れの記録から入れる", BOLD),
        ("   ・販売手数料・FBA配送手数料・保管手数料・返金額：日付範囲別レポートの各合計（マイナス表示ならプラスに直して入れる）", BLACK),
        ("   ・1個あたり着地原価：工場への支払い＋国際送料＋関税＋通関・検品・国内配送 を、その仕入れの個数で割った額", BLACK),
        ("   ・返品個数のうち、売り物にならなかった個数を「再販不可」に入れる（再販できた分は在庫に戻るので原価から外れる）", BLACK),
        ("   ・広告費は、Amazonの入金から引かれていても、カードで払っていても、ここに入れる", BLACK),
        ("   ・Amazon以外の支払い：カードで払った広告費・倉庫への送料など、入金から引かれていない支払い（現金の増減の計算に使う）", BLACK),
        ("2. 判定：赤字＝利益がマイナス／注意＝広告前利益率・返品率・在庫月数のどれかが設定の基準を外れている", BOLD),
        ("3. 内訳シート：期間全体で、売上がどこに消えたかを、大きい順に並べる", BOLD),
        ("4. 改善シミュレーション：原価・価格・広告・返品・FBA手数料を変えると、1個あたりの利益がどう変わるかを見る", BOLD),
        ("", BLACK),
        ("見る順番", BOLD),
        ("① 1個あたり利益がマイナスの月はないか → ② 広告前利益率が薄すぎないか（薄いと広告と返品で赤字になる） → ③ 内訳の1位はどこか → ④ 在庫月数と現金の増減", BLACK),
        ("", BLACK),
        ("注意：この表の利益は、事務所・ツール代などの固定費と税金を引く前の数字です。消費税の扱いは、自分の申告方法に合わせて読み替えてください。", BLACK),
    ]
    for i, (t, f) in enumerate(lines, start=1):
        ws.cell(i, 1, t).font = f
    ws.column_dimensions["A"].width = 130


def build_settings(wb):
    ws = wb.create_sheet("設定")
    ws["A1"], ws["A1"].font = "判定の基準（初期値。自分の商品に合わせて変える）", TITLE
    rows = [
        ("広告前利益率の下限", 0.10, PCT, "これより薄いと、広告と返品で赤字になりやすい"),
        ("返品率の上限", 0.05, PCT, "これを超えたら、不良・説明文・梱包を点検する"),
        ("在庫月数の上限", 3.0, MONTHS, "今の販売ペースで何か月分の在庫か。超えると保管料と資金が寝る"),
    ]
    for i, (label, v, fmt, note) in enumerate(rows, start=4):
        ws.cell(i, 2, label).font = BOLD
        c = ws.cell(i, 3, v)
        c.font, c.fill, c.number_format, c.border = BLUE, INPUT, fmt, BOX
        ws.cell(i, 4, note).font = BLACK
    ws.column_dimensions["B"].width, ws.column_dimensions["C"].width, ws.column_dimensions["D"].width = 22, 10, 70


def build_monthly(wb):
    ws = wb.create_sheet("月次")
    ws["A1"], ws["A1"].font = "月次の監査", TITLE
    ws["A2"] = "緑の見出し＝入力／灰色の見出し＝自動計算。1か月1行。【見本】の6か月分は架空の数字です。"
    ws["A3"] = "実質売上＝売上−返金額／原価＝（販売個数−再販できた返品）×着地原価／利益＝実質売上−原価−Amazon費用−広告費"
    for c in ("A2", "A3"):
        ws[c].font = BLACK
    hr = FIRST - 1
    for j, (name, is_in, width, fmt, formula) in enumerate(COLS, start=1):
        head(ws, hr, j, name, is_in)
        ws.column_dimensions[ws.cell(hr, j).column_letter].width = width
        for r in range(FIRST, LAST + 1):
            c = ws.cell(r, j)
            c.border = BOX
            if fmt:
                c.number_format = fmt
            if formula:
                c.value, c.font = formula.format(r=r), BLACK
            else:
                c.font, c.fill = BLUE, INPUT
    ws.row_dimensions[hr].height = 40
    for i, row in enumerate(SAMPLE):
        for j, v in enumerate(row, start=1):
            ws.cell(FIRST + i, j, v)
    rng = f"AD{FIRST}:AD{LAST}"
    for word, fill in (("赤字", RED_FILL), ("注意", YEL_FILL), ("OK", OK_FILL)):
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'$AD{FIRST}="{word}"'], fill=fill))
    ws.conditional_formatting.add(f"V{FIRST}:V{LAST}", FormulaRule(formula=[f"AND(ISNUMBER($V{FIRST}),$V{FIRST}<0)"], fill=RED_FILL))
    ws.freeze_panes = ws.cell(FIRST, 2)


def build_breakdown(wb):
    ws = wb.create_sheet("内訳")
    ws["A1"], ws["A1"].font = "売上はどこに消えたか（月次シートの全期間）", TITLE
    m = "'月次'!"
    rng = lambda col: f"{m}${col}${FIRST}:${col}${LAST}"
    ws["A3"], ws["A3"].font = "期間の合計", BOLD
    totals = [
        ("売上（税込）", f"=SUM({rng('B')})", YEN),
        ("販売個数", f"=SUM({rng('C')})", NUM),
        ("返品個数", f"=SUM({rng('J')})", NUM),
        ("実質の販売個数", "=C5-C6", NUM),
        ("利益", f"=SUM({rng('U')})", YEN),
        ("1個あたり利益", '=IF(C7>0,C8/C7,0)', YEN),
        ("広告前利益", f"=SUM({rng('T')})", YEN),
        ("広告前利益率", f'=IFERROR(C10/SUM({rng("Q")}),0)', PCT),
        ("現金の増減", f"=SUM({rng('AC')})", YEN),
    ]
    for i, (label, f, fmt) in enumerate(totals, start=4):
        ws.cell(i, 2, label).font = BOLD
        c = ws.cell(i, 3, f)
        c.font, c.number_format, c.border = BLACK, fmt, BOX

    ws["A15"], ws["A15"].font = "売上から出ていったお金（売上に対する割合）", BOLD
    for j, h in enumerate(["項目", "金額", "売上比", "1個あたり", "順位"], start=2):
        head(ws, 16, j, h, False)
    items = [
        ("商品原価", rng("R")),
        ("販売手数料", rng("D")),
        ("FBA配送手数料", rng("E")),
        ("在庫保管手数料", rng("F")),
        ("その他のAmazon費用", rng("G")),
        ("広告費", rng("H")),
        ("返金", rng("I")),
    ]
    first, last = 17, 17 + len(items) - 1
    for k, (label, r) in enumerate(items):
        row = first + k
        ws.cell(row, 2, label).font = BLACK
        cells = [
            (f"=SUM({r})", YEN),
            (f"=IF($C$4>0,C{row}/$C$4,0)", PCT),
            (f"=IF($C$7>0,C{row}/$C$7,0)", YEN),
            (f"=RANK(C{row},$C${first}:$C${last})", "0"),
        ]
        for j, (f, fmt) in enumerate(cells, start=3):
            c = ws.cell(row, j, f)
            c.font, c.number_format, c.border = BLACK, fmt, BOX
        ws.cell(row, 2).border = BOX
    row = last + 1
    ws.cell(row, 2, "残った利益").font = BOLD
    for j, (f, fmt) in enumerate([("=C8", YEN), ("=IF($C$4>0,C8/$C$4,0)", PCT), ("=C9", YEN)], start=3):
        c = ws.cell(row, j, f)
        c.font, c.number_format, c.border = BOLD, fmt, BOX
    ws.cell(row + 2, 2, "順位1〜2の項目から手を付ける。広告は最後（広告前の利益が薄いまま広告を止めても、売上が落ちるだけで赤字は残りやすい）").font = BLACK
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 22
    for col in "CDEF":
        ws.column_dimensions[col].width = 14


def build_sim(wb):
    ws = wb.create_sheet("改善シミュレーション")
    ws["A1"], ws["A1"].font = "改善シミュレーション（1個あたり。内訳シートの期間平均が基準）", TITLE
    ws["A2"] = "販売個数は変わらないものとして計算します。値上げで販売個数が減る影響は入っていません。"
    ws["A2"].font = BLACK
    b = "'内訳'!"
    ws["B4"], ws["B4"].font = "変える条件", BOLD
    levers = [
        ("販売価格", 0.0, "＋なら値上げ。販売手数料・返金も同じ割合で変わる"),
        ("着地原価", -0.05, "工場との交渉・発注ロット・輸送の見直し"),
        ("広告費", 0.0, "−なら削減"),
        ("返品・返金", -0.30, "不良・説明文・梱包の改善"),
        ("FBA配送手数料", 0.0, "梱包を小さくしてサイズ区分を下げるなど"),
    ]
    for i, (label, v, note) in enumerate(levers, start=5):
        ws.cell(i, 2, label).font = BOLD
        c = ws.cell(i, 3, v)
        c.font, c.fill, c.number_format, c.border = BLUE, INPUT, '+0%;-0%;0%', BOX
        ws.cell(i, 4, note).font = BLACK

    for j, h in enumerate(["項目（1個あたり）", "いま", "変えた後", "差"], start=2):
        head(ws, 11, j, h, False)
    per = lambda r: f"IF({b}$C$7>0,{b}$C${r}/{b}$C$7,0)"
    sales_now = f"=IF({b}$C$7>0,SUM('月次'!$B${FIRST}:$B${LAST})/{b}$C$7,0)"
    rows = [
        ("売上", sales_now, "=C12*(1+$C$5)"),
        ("返金", f"={per(23)}", "=C13*(1+$C$5)*(1+$C$8)"),
        ("商品原価", f"={per(17)}", "=C14*(1+$C$6)"),
        ("販売手数料", f"={per(18)}", "=C15*(1+$C$5)"),
        ("FBA配送手数料", f"={per(19)}", "=C16*(1+$C$9)"),
        ("保管・その他のAmazon費用", f"={per(20)}+{per(21)}", "=C17"),
        ("広告費", f"={per(22)}", "=C18*(1+$C$7)"),
    ]
    for k, (label, now, after) in enumerate(rows):
        r = 12 + k
        ws.cell(r, 2, label).font = BLACK
        ws.cell(r, 2).border = BOX
        for j, f in ((3, now), (4, after), (5, f"=D{r}-C{r}")):
            c = ws.cell(r, j, f)
            c.font, c.number_format, c.border = BLACK, YEN, BOX
    r = 19
    ws.cell(r, 2, "1個あたり利益").font = BOLD
    for j, col in ((3, "C"), (4, "D")):
        c = ws.cell(r, j, f"={col}12-SUM({col}13:{col}18)")
        c.font, c.number_format, c.border = BOLD, YEN, BOX
    c = ws.cell(r, 5, "=D19-C19")
    c.font, c.number_format, c.border = BOLD, YEN, BOX
    ws.cell(20, 2, "期間の利益（実質の販売個数×1個あたり）").font = BOLD
    for j, col in ((3, "C"), (4, "D")):
        c = ws.cell(20, j, f"={col}19*{b}$C$7")
        c.font, c.number_format, c.border = BOLD, YEN, BOX
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 34
    for col in "CDE":
        ws.column_dimensions[col].width = 14


def main():
    wb = Workbook()
    build_guide(wb)
    build_settings(wb)
    build_monthly(wb)
    build_breakdown(wb)
    build_sim(wb)
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
