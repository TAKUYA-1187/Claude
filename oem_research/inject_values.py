# -*- coding: utf-8 -*-
"""Write cached formula results into an openpyxl-built workbook without letting
LibreOffice rewrite (and restyle) it.

usage: python3 inject_values.py styled.xlsx computed.xlsx out.xlsx
  styled.xlsx   : workbook written by openpyxl (formulas, no cached values)
  computed.xlsx : the same workbook after LibreOffice recalculation
"""
import sys
import zipfile

from lxml import etree
from openpyxl import load_workbook

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
RNS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PNS = "http://schemas.openxmlformats.org/package/2006/relationships"
ERRORS = {"#DIV/0!", "#N/A", "#NAME?", "#NULL!", "#NUM!", "#REF!", "#VALUE!"}


def sheet_paths(zf):
    wb = etree.fromstring(zf.read("xl/workbook.xml"))
    rels = etree.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
    target = {r.get("Id"): r.get("Target") for r in rels.findall(f"{{{PNS}}}Relationship")}
    out = {}
    for s in wb.find(f"{{{NS}}}sheets"):
        t = target[s.get(f"{{{RNS}}}id")].lstrip("/")
        out[t if t.startswith("xl/") else "xl/" + t] = s.get("name")
    return out


def main(styled, computed, out):
    values = load_workbook(computed, data_only=True)
    filled = missing = 0
    with zipfile.ZipFile(styled) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        paths = sheet_paths(zin)
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename in paths:
                ws = values[paths[item.filename]]
                root = etree.fromstring(data)
                for c in root.iter(f"{{{NS}}}c"):
                    f = c.find(f"{{{NS}}}f")
                    if f is None:
                        continue
                    val = ws[c.get("r")].value
                    v = c.find(f"{{{NS}}}v")
                    if v is None:
                        v = etree.SubElement(c, f"{{{NS}}}v")
                    if val is None:
                        missing += 1
                        c.remove(v)
                        continue
                    if isinstance(val, bool):
                        c.set("t", "b")
                        v.text = "1" if val else "0"
                    elif isinstance(val, (int, float)):
                        if "t" in c.attrib:
                            del c.attrib["t"]
                        v.text = repr(float(val)) if isinstance(val, float) else str(val)
                    elif isinstance(val, str) and val in ERRORS:
                        c.set("t", "e")
                        v.text = val
                    else:
                        c.set("t", "str")
                        v.text = str(val)
                    filled += 1
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            zout.writestr(item, data)
    print(f"cached values written: {filled}, formulas without a value: {missing}")


if __name__ == "__main__":
    main(*sys.argv[1:4])
