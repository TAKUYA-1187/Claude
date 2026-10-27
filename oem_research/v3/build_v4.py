# -*- coding: utf-8 -*-
"""Build the v4 workbook: every Amazon listing >=10,000 yen selling >=100/month in 231
large/heavy keywords, judged with 天瞳's published rates, on the user's 天瞳 template.

usage: OEM_DATA=<dir with scan_v4/prods_v4/detail_v4/cn_v4/cn_fin/eval_final .json> \
       python3 build_v4.py <template.xlsx> <out.xlsx>

Finalist sheets are copies of the template (formulas untouched) with both rival blocks
filled from the Amazon pages, plus a plan block that recomputes our product with the
template formula and the real P&L. Every formula reads the 前提条件 sheet (天瞳 rates).
"""
import json
import os
import re
import statistics as st
import sys
from urllib.parse import quote

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "v2"))
import analyze_prods as P_  # noqa: E402
import analyze_v4 as A  # noqa: E402
import build_v2 as B  # noqa: E402
import complaints as C  # noqa: E402
import importlib  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from openpyxl.formatting.rule import CellIsRule, FormulaRule  # noqa: E402
from openpyxl.styles import Alignment, Font  # noqa: E402
from finalists_cn import T2  # noqa: E402
from type_map import KW2T  # noqa: E402

D = importlib.import_module(os.environ.get("OEM_DATAMOD", "data_v4"))

put, head_row, merge_put, link_sheet = B.put, B.head_row, B.merge_put, B.link_sheet
YEN, PCT, CNY, KG, CM = B.YEN, B.PCT, B.CNY, B.KG, B.CM
M3 = '0.000"m³"'
AMZ = "https://www.amazon.co.jp/dp/"
AS = "'前提条件'"
OV = "'01_最終候補'"
P = {k: f"{AS}!$B${r}" for k, r in dict(
    fx=4, tfee=5, tvat=6, tduty=7, fr_t=8, fr_s=9, fr_b=10, fr_40=11,
    feeA=13, feeB=14, fth=15, ad10=16, ad5=17, ret=18, stor=19, sdays=20, cn=21, tax=22, fr=23,
    go_t=26, go_r=27, c_t=28, c_r=29).items()}
SG = dict(up=f"{AS}!$B$34:$B$44", kg=f"{AS}!$C$34:$C$44", rate=f"{AS}!$D$34:$D$44")


def ship_f(L, W, H, kg, n):
    """Sagawa rate for n parcels (longest side split), smallest rate whose size and weight both fit."""
    return (f"={n}*_xlfn.MINIFS({SG['rate']},{SG['up']},\">=\"&({L}/{n}+{W}+{H}),{SG['kg']},\">=\"&({kg}/{n}))")


def fee_f(price, cls):
    return (f"IF({cls}=\"F\",{P['feeA']}*MIN({price},{P['fth']})+{P['feeB']}*MAX(0,{price}-{P['fth']}),"
            f"{price}*IF({cls}=\"A\",{P['feeA']},{P['feeB']}))")


def load(name):
    return json.load(open(os.path.join(os.environ["OEM_DATA"], name), encoding="utf-8"))


SCAN, PRODS, DET = load("scan_v4.json"), load("prods_v4.json"), load("detail_v4.json")
CN, FIN, EV = load("cn_v4.json"), load("cn_fin.json"), load("eval_final.json")
EVA = {r["asin"]: r for r in EV}


# ---------------------------------------------------------------------------
def build_assumptions(ws):
    ws.title = "前提条件"
    ws.column_dimensions["A"].width = 46
    for col, w in zip("BCDEFGH", (14, 12, 12, 60, 3, 12, 12)):
        ws.column_dimensions[col].width = w
    put(ws, "A1", "前提条件（天瞳 送料・保管料一覧表 2024.1.6版／2024年8月1日以降の料金）", B.F_TITLE, border=False)
    put(ws, "A2", "黄色セル＝入力値。ここを変えると01_最終候補・各R0xシート・ロードマップが再計算されます（02〜04シートは同じ式で計算した値）。",
        B.F_NOTE, border=False)
    put(ws, "A3", "① 天瞳リサーチ表（テンプレート）の計算パラメータ・国際送料", B.F_SUB, border=False)
    rows = [
        (4, "人民元レート（円/元）", 23.5, "0.00", "テンプレートI7と同じ"),
        (5, "Amazon手数料（テンプレート）", 0.165, PCT, "テンプレートL7と同じ16.5%"),
        (6, "消費税（輸入・仕入原価×）", 0.10, PCT, "テンプレートM7と同じ"),
        (7, "関税（仕入原価×）", 0.10, PCT, "テンプレートN7と同じ（実際は家具・金属製品の多くが0%）"),
        (8, "海上運賃 2〜20m³（円/m³）", 17325, YEN, "料金表：2m³以上20m³未満。テンプレートO7と同じ"),
        (9, "海上運賃 0〜2m³（円/m³）", 33000, YEN, "料金表：2m³未満（初回テスト発注）"),
        (10, "海上運賃 20m³以上（円/m³）", 13125, YEN, "料金表：20m³以上。大型品を1品50〜100個ずつ発注すると20m³を超える"),
        (11, "40フィートコンテナ（円/m³換算）", 10450, YEN, "料金表：40ft 1本（約67m³）。年商5億の段階"),
    ]
    rows2 = [
        (13, "販売手数料A（ホーム・DIY・ペット・ベビー・オフィス等）", 0.154, PCT, "2026年改定後（税抜）。家具は下の区分F"),
        (14, "販売手数料B（スポーツ・アウトドア・カー用品）", 0.104, PCT, "同上"),
        (15, "家具（区分F）：15.4%が適用される上限額", 20000, YEN, "家具は売価2万円までの部分15.4%、2万円を超える部分10.4%"),
        (16, "広告費（立ち上げ期）売上比", 0.10, PCT, "1位獲得までの期間"),
        (17, "広告費（1位定着後）売上比", 0.05, PCT, "オーガニック中心になった後"),
        (18, "返品・破損 売上比", 0.04, PCT, "大型品は配送破損を見込む"),
        (19, "保管料（円/m³/日）", 77, YEN, "料金表：1m³あたり1日77円（SKUごと切上げ）"),
        (20, "平均保管日数", 45, "0", "入庫から出荷までの平均。発注サイクル1.5か月"),
        (21, "中国側諸費用（検品・国内運送・代行）仕入原価比", 0.08, PCT, ""),
        (22, "消費税率（売上）", 0.10, PCT, "売上の消費税は納税。輸入消費税は仕入税額控除で戻る"),
        (23, "国際送料（実質試算で使う単価）", f"={P['fr_b']}", YEN, "20m³以上の単価。初年度の小ロットはB8（17,325円）"),
    ]
    for r, label, v, fmt, note in rows + rows2:
        put(ws, f"A{r}", label)
        put(ws, f"B{r}", v, fmt=fmt, fill=B.FILL_IN)
        put(ws, f"C{r}", note, B.F_NOTE, align=B.LEFT_C, border=False)
    put(ws, "A12", "② 実質試算（Claude）の追加パラメータ", B.F_SUB, border=False)
    put(ws, "A25", "③ 判定基準", B.F_SUB, border=False)
    for r, label, v in ((26, "GO：テンプレ利益率 下限", 0.30), (27, "GO：実質利益率（広告5%）下限", 0.15),
                        (28, "COND：テンプレ利益率 下限", 0.25), (29, "COND：実質利益率（広告5%）下限", 0.10)):
        put(ws, f"A{r}", label)
        put(ws, f"B{r}", v, fmt=PCT, fill=B.FILL_IN)
    put(ws, "C26", "両方を満たせばGO。テンプレ利益率は広告・返品・保管・売上消費税を含まないため実質より15〜20pt高く出る。",
        B.F_NOTE, align=B.LEFT_C, border=False)
    put(ws, "A31", "④ 国内送料（佐川急便・天瞳料金表。梱包資材・出荷作業込み）", B.F_SUB, border=False)
    for col, h in zip("ABCD", ("サイズ", "三辺合計（以下）cm", "重量（以下）kg", "送料（円/個口）")):
        put(ws, f"{col}33", h, B.F_HEAD, fill=B.FILL_HEAD, align=B.WRAP_C)
    names = ["60", "80", "100", "140（120は表に無く140扱い）", "160", "170", "180", "200", "220", "240", "260"]
    for i, ((up, kg, rate), nm) in enumerate(zip(A.SAGAWA, names)):
        r = 34 + i
        put(ws, f"A{r}", nm + "サイズ")
        put(ws, f"B{r}", up)
        put(ws, f"C{r}", kg)
        put(ws, f"D{r}", rate, fmt=YEN, fill=B.FILL_IN)
    put(ws, "A46", "260cm超・50kg超は2〜3個口に分割（料金表「特大サイズ」）。送料＝個口数×各個口の送料。サンプル航空便1,500円/kg、サンプル海上便15,000円/m³。",
        B.F_NOTE, border=False)
    put(ws, "A48", "⑤ 読み方・注意", B.F_SUB, border=False)
    notes = [
        "・テンプレ利益率（天瞳リサーチ表と同じ）＝売価−原価−国内送料−手数料16.5%−輸入消費税−関税−国際送料（17,325円/m³）。",
        "・実質利益率＝税抜売上−販売手数料−国内送料−返品−保管（m³×77円×45日）−（原価×1.08＋付加価値費＋国際送料13,125円/m³）−広告費。",
        "・天瞳の国内送料は梱包資材・出荷作業込みのため、出荷作業費は別計上していない。",
        "・月販は商品ページの『過去1か月でN点以上購入』バッジの下限値（実数はこれ以上）。",
        "・1688原価：最終候補（01シート）は仕様が一致する1688出品をタイトルで抽出して決めた値（各R0xシート下部に根拠の出品）。一次スクリーニング（02シート）は同タイプの1688出品の中間50%の中央値（標準原価）。掲載価格は少量の価格で、50個以上の発注では通常1〜2割下がる。",
    ]
    for i, t in enumerate(notes):
        r = 49 + i
        ws.merge_cells(f"A{r}:H{r}")
        put(ws, f"A{r}", t, B.F_SMALL, align=B.WRAP_TL, border=False)
        ws.row_dimensions[r].height = 28
    ws.freeze_panes = "A4"


# ---------------------------------------------------------------------------
OV_HEAD = ["No", "企画品（自社仕様）", "タイプ", "手数料区分", "定常売価", "1688原価（元）", "付加価値費",
           "梱包 長", "梱包 幅", "梱包 高", "梱包重量", "個口", "CBM", "三辺合計", "国内送料",
           "テンプレ粗利", "テンプレ利益率", "税抜売上", "販売手数料", "保管料", "着地原価",
           "実質利益（広告10%）", "実質利益（広告5%）", "実質利益率（広告5%）", "判定",
           "価格差（倍）", "競合1位の月販", "月販目標（3年目）", "月利（3年目）", "詳細シート", "規制・注意",
           "Amazon（競合1位）", "1688（原価の根拠・月销最多）", "Alibaba検索"]


def build_overview(ws):
    ws.title = "01_最終候補"
    put(ws, "A1", "最終候補の商品化判断（天瞳の実料金・数式で再計算）", B.F_TITLE, border=False)
    put(ws, "A2", "黄色＝入力（売価・原価・梱包）。青字＝推定値（要実測・要見積）。送料・利益はすべて前提条件シートから計算。", B.F_NOTE, border=False)
    widths = [5, 40, 10, 6, 10, 9, 8, 6, 6, 6, 7, 5, 8, 7, 9, 10, 8, 10, 9, 8, 10, 10, 10, 9, 7, 7, 8, 8, 12, 14, 24, 34, 40, 30]
    head_row(ws, 4, OV_HEAD, widths, height=40)
    rowmap = {}
    for i, f in enumerate(D.FINALISTS + D.REJECTED):
        r = 5 + i
        rowmap[f["id"]] = r
        L, W, H, kg = f["pkg"]
        vals = {
            "A": f["id"], "B": f["name"], "C": f["type_jp"], "D": f["fee"], "E": f["price"], "F": round(f["cny"], 1), "G": f["va"],
            "H": L, "I": W, "J": H, "K": kg, "L": f["parcels"],
            "M": f"=H{r}*I{r}*J{r}/1000000", "N": f"=H{r}+I{r}+J{r}", "O": ship_f(f"H{r}", f"I{r}", f"J{r}", f"K{r}", f"L{r}"),
            "P": f"=E{r}-F{r}*{P['fx']}-O{r}-E{r}*{P['tfee']}-F{r}*{P['fx']}*({P['tvat']}+{P['tduty']})-{P['fr_t']}*M{r}",
            "Q": f"=P{r}/E{r}", "R": f"=E{r}/(1+{P['tax']})", "S": "=" + fee_f(f"E{r}", f"D{r}"),
            "T": f"=M{r}*{P['stor']}*{P['sdays']}", "U": f"=F{r}*{P['fx']}*(1+{P['cn']})+G{r}+{P['fr']}*M{r}",
            "V": f"=R{r}-S{r}-O{r}-R{r}*{P['ret']}-T{r}-U{r}-R{r}*{P['ad10']}",
            "W": f"=R{r}-S{r}-O{r}-R{r}*{P['ret']}-T{r}-U{r}-R{r}*{P['ad5']}", "X": f"=W{r}/R{r}",
            "Y": f"=IF(AND(Q{r}>={P['go_t']},X{r}>={P['go_r']}),\"GO\",IF(AND(Q{r}>={P['c_t']},X{r}>={P['c_r']}),\"COND\",\"NO\"))",
            "Z": f"=E{r}/(F{r}*{P['fx']})", "AA": f["top_units"], "AB": f["y"][2] if "y" in f else 0, "AC": f"=AB{r}*W{r}",
            "AE": f["reg"],
        }
        fmts = {"E": YEN, "F": CNY, "G": YEN, "H": CM, "I": CM, "J": CM, "K": KG, "M": M3, "N": CM, "O": YEN, "P": YEN,
                "Q": PCT, "R": YEN, "S": YEN, "T": YEN, "U": YEN, "V": YEN, "W": YEN, "X": PCT, "Z": '0.0"倍"',
                "AA": "#,##0", "AB": "#,##0", "AC": YEN}
        for col, v in vals.items():
            est = True
            fill = B.FILL_IN if col in ("E", "F", "G", "H", "I", "J", "K", "L", "AB") else None
            font = B.F_EST if est and col in ("F", "H", "I", "J", "K") else (B.F_BOLD if col in ("X", "Y", "AC") else B.F_SMALL)
            put(ws, f"{col}{r}", v, font, fmt=fmts.get(col), fill=fill,
                align=B.LEFT_C if col in ("B", "AE") else B.WRAP_C)
        if f.get("sheet"):
            cell = put(ws, f"AD{r}", f["sheet"], B.F_LINK)
            link_sheet(cell, f"'{f['sheet']}'!A1")
        else:
            put(ws, f"AD{r}", "（検証で見送り）", B.F_SMALL, align=B.WRAP_C)
        for col, url in links_of(f):
            cell = put(ws, f"{col}{r}", url, B.F_LINK, align=B.LEFT_C)
            cell.hyperlink = url
        ws.row_dimensions[r].height = 32
    last = 4 + len(D.FINALISTS) + len(D.REJECTED)
    ws.freeze_panes = "C5"
    for txt, fill in (("GO", B.FILL_GO), ("COND", B.FILL_COND), ("NO", B.FILL_NO)):
        ws.conditional_formatting.add(f"Y5:Y{last}", CellIsRule(operator="equal", formula=[f'"{txt}"'], fill=fill))
    r = last + 2
    put(ws, f"A{r}", "注：判定は前提条件③（テンプレ30%以上かつ実質15%以上でGO）。月利＝月販目標×実質利益（広告5%）。価格差＝売価÷（1688原価×レート）。"
                  f"X行（{len(D.REJECTED)}件）は上位の売れ筋だが、1688の仕様一致品の価格で検証して見送った商品。", B.F_NOTE, border=False)
    ws.row_dimensions[5 + len(D.FINALISTS)].height = 32
    return rowmap, last


def links_of(f):
    """(column, url) for a finalist: rival 1 on Amazon, best-selling spec-matched 1688 offer, Alibaba search."""
    out = []
    if f.get("rivals"):
        out.append(("AF", AMZ + f["rivals"][0]))
    if f.get("fin_key"):
        ev = evidence(f, 1)
        if ev and ev[0].get("url"):
            out.append(("AG", ev[0]["url"]))
    if f.get("en"):
        out.append(("AH", "https://www.alibaba.com/trade/search?SearchText=" + quote(f["en"])))
    return out


def ref_offer(x):
    """Best-selling 1688 offer of the product's type priced near its mainstream cost (0.5-2x)."""
    offers = CN.get(x["type"]) or (SCAN.get(x["kw"], {}).get("offers") or [])
    c = x.get("cost")
    if c:
        offers = [o for o in offers if 0.5 * c <= o["price"] <= 2 * c]
    offers = [o for o in offers if o.get("url")]
    return max(offers, key=lambda o: o["sold"]) if offers else None


# ---------------------------------------------------------------------------
SCREEN_HEAD = ["全条件", "判定", "タイプ", "ASIN", "商品名", "売価", "月販（下限）", "評価数", "★", "サブカテゴリ順位",
               "梱包 長×幅×高", "梱包重量", "梱包の根拠", "160サイズ以上", "1688標準原価（元）", "1688同タイプ出品数", "価格差（倍）",
               "国内送料", "個口", "CBM", "テンプレ利益率", "実質利益率（広告5%）", "1個利益（広告5%）", "月利見込（月販×1個利益）",
               "ブランド依存", "規制・注意", "検索キーワード", "Amazon商品ページ", "1688参考出品（同タイプ・月销最多）",
               "1688参考 価格（元）", "Alibaba検索"]


def build_screen(ws):
    ws.title = "02_全商品スクリーニング"
    put(ws, "A1", f"Amazon全商品スクリーニング：{len(SCAN)}キーワードで¥10,000以上・月{D.SCREEN_MIN_UNITS}個以上売れている{len(EV)}商品",
        B.F_TITLE, border=False)
    put(ws, "A2", "全条件＝売価¥1万以上・月販100個以上・160サイズ以上（三辺160cm超または30kg以上）・1688との価格差3倍以上・判定GO/COND・ブランド依存でない。"
                  "梱包は商品ページ記載値、無い場合は同種品の標準梱包（青字の推定）。原価は同タイプ1688出品の中間50%の中央値（標準原価）。最終候補は01シートで仕様一致の原価に置換。",
        B.F_NOTE, border=False)
    head_row(ws, 4, SCREEN_HEAD, [12, 6, 10, 12, 50, 9, 8, 8, 5, 22, 14, 7, 14, 6, 9, 7, 7, 8, 5, 7, 8, 8, 9, 11, 6, 16, 18, 34, 40, 9, 30],
             height=40)
    order = {"GO": 0, "COND": 1, "NO": 2, None: 3}
    vno = getattr(D, "VERIFIED_NO", {})
    rows = sorted(EV, key=lambda x: (not x["ok_all"], bool(vno.get(x["type"]) or vno.get(x["kw"])), order[x.get("judge")],
                                     -(x.get("profit_month") or -1e9), -x["bought"]))
    tj = {v["id"]: v["jp"] for v in D.TYPES}
    for i, x in enumerate(rows):
        r = 5 + i
        vno = getattr(D, "VERIFIED_NO", {})
        vkey = vno.get(x["type"]) or vno.get(x["kw"])
        flag = ("×（検証で見送り " + vkey + "）" if vkey else "○") if x["ok_all"] else ""
        vals = [flag, x.get("judge") or "原価不明", tj.get(x["type"], x["type"] or x["kw"]), x["asin"], x["title"][:90],
                x["price"], x["bought"], x["reviews"], x["rating"],
                f"{x['sub_cat']} {x['sub_rank']:,}位" if x.get("sub_rank") else "",
                f"{x['L']:.0f}×{x['W']:.0f}×{x['H']:.0f}" if x.get("L") else "", x.get("kg"), (x.get("pk_src") or "")[:14],
                ("○" if x.get("size160") else "×") if x.get("L") else "", x.get("cost"), x.get("n_offers"), x.get("gap"),
                x.get("dom"), x.get("parcels"), x.get("cbm"), x.get("tm"), x.get("r5"), x.get("p5"), x.get("profit_month"),
                "○" if x["brand_lock"] else "", x.get("reg", ""), x["kw"]]
        ro = ref_offer(x)
        en = D.TYPE_EN.get(x["type"]) if hasattr(D, "TYPE_EN") else None
        vals += [AMZ + x["asin"], ro["url"] if ro else "", ro["price"] if ro else None,
                 "https://www.alibaba.com/trade/search?SearchText=" + quote(en) if en else ""]
        fmts = [None, None, None, None, None, YEN, "#,##0", "#,##0", "0.0", None, None, KG, None, None, "#,##0.0", "#,##0",
                '0.0"倍"', YEN, "0", M3, PCT, PCT, YEN, YEN, None, None, None, None, None, "#,##0.0", None]
        for j, (v, fm) in enumerate(zip(vals, fmts), start=1):
            est = j in (11, 12) and "推定" in (x.get("pk_src") or "")
            cell = put(ws, f"{B.get_column_letter(j)}{r}", v, B.F_EST if est else B.F_SMALL, fmt=fm,
                       align=B.LEFT_C if j == 5 else None)
            if j == 4 or (j in (28, 29, 31) and v):
                cell.hyperlink = v if j != 4 else AMZ + x["asin"]
                cell.font = B.F_LINK
    last = 4 + len(rows)
    ws.freeze_panes = "F5"
    ws.auto_filter.ref = f"A4:{B.get_column_letter(len(SCREEN_HEAD))}{last}"
    ws.conditional_formatting.add(f"A5:A{last}", CellIsRule(operator="equal", formula=['"○"'], fill=B.FILL_GO))
    for txt, fill in (("GO", B.FILL_GO), ("COND", B.FILL_COND)):
        ws.conditional_formatting.add(f"B5:B{last}", CellIsRule(operator="equal", formula=[f'"{txt}"'], fill=fill))


# ---------------------------------------------------------------------------
def type_stats():
    out = []
    for t in D.TYPES:
        xs = [x for x in EV if x["type"] == t["id"]]
        if not xs:
            continue
        top = max(xs, key=lambda x: (x["bought"], x["reviews"]))
        ok = [x for x in xs if x["ok_all"]]
        gaps = [x["gap"] for x in xs if x.get("gap")]
        offers = CN.get(t["id"], [])
        out.append(dict(t=t, n=len(xs), units=sum(x["bought"] for x in xs), sales=sum(x["bought"] * x["price"] for x in xs),
                        med_price=st.median([x["price"] for x in xs]), top=top, n_off=len(offers),
                        cn_sold=sum(o["sold"] for o in offers), cn_med=st.median([o["price"] for o in offers]) if offers else None,
                        gap=st.median(gaps) if gaps else None,
                        s160=sum(1 for x in xs if x.get("size160")) / len(xs),
                        go=sum(1 for x in xs if x.get("judge") == "GO"), ok=len(ok),
                        best=max([x["profit_month"] for x in ok], default=None)))
    out.sort(key=lambda s: -s["sales"])
    return out


def build_types(ws):
    ws.title = "03_タイプ別集計"
    put(ws, "A1", "商品タイプ別の集計（Amazon上位商品 × 1688同タイプ出品）", B.F_TITLE, border=False)
    put(ws, "A2", "Amazon：¥1万以上・月100個以上の商品のみ。1688：タイプ別に3〜5語で検索し、タイトルが合う出品のみ（部品・付属品は価格下限で除外）。",
        B.F_NOTE, border=False)
    head_row(ws, 4, ["タイプ", "対象商品数", "月販合計（下限）", "月商合計（下限）", "中央売価", "販売数1位の商品", "1位の売価", "1位の月販",
                     "1688出品数", "1688月销合計", "1688中央値（元）", "価格差 中央値", "160サイズ以上の割合", "GO商品数", "全条件を満たす商品数",
                     "最良の月利見込", "結論", "理由・注意"],
             [14, 7, 9, 12, 9, 46, 9, 7, 7, 9, 9, 8, 8, 6, 8, 11, 12, 50], height=40)
    for i, s in enumerate(type_stats()):
        r = 5 + i
        t, top = s["t"], s["top"]
        vals = [t["jp"], s["n"], s["units"], s["sales"], s["med_price"], top["title"][:70], top["price"], top["bought"],
                s["n_off"], s["cn_sold"], s["cn_med"], s["gap"], s["s160"], s["go"], s["ok"], s["best"], t["verdict"], t["why"]]
        fmts = [None, "0", "#,##0", YEN, YEN, None, YEN, "#,##0", "0", "#,##0", "#,##0.0", '0.0"倍"', "0%", "0", "0", YEN, None, None]
        for j, (v, fm) in enumerate(zip(vals, fmts), start=1):
            cell = put(ws, f"{B.get_column_letter(j)}{r}", v, B.F_BOLD if j == 17 else B.F_SMALL, fmt=fm,
                       align=B.WRAP_TL if j in (6, 18) else B.WRAP_C)
            if j == 6:
                cell.hyperlink = AMZ + top["asin"]
                cell.font = B.F_LINK
        ws.row_dimensions[r].height = 36
    last = 4 + len(type_stats())
    ws.freeze_panes = "B5"
    ws.auto_filter.ref = f"A4:R{last}"
    for txt, fill in (("採用", B.FILL_GO), ("条件付き", B.FILL_COND), ("見送り", B.FILL_NO)):
        ws.conditional_formatting.add(f"Q5:Q{last}", FormulaRule(formula=[f'LEFT(Q5,{len(txt)})="{txt}"'], fill=fill))


# ---------------------------------------------------------------------------
def build_dresser(ws):
    ws.title = "04_ドレッサー検証"
    put(ws, "A1", "ドレッサー・鏡台の検証（ご指摘の商品群）", B.F_TITLE, border=False)
    put(ws, "A2", D.DRESSER["lead"], B.F_BASE, border=False)
    ws.merge_cells("A2:N2")
    ws["A2"].alignment = B.WRAP_TL
    ws.row_dimensions[2].height = 60
    head_row(ws, 4, ["ASIN", "商品名", "売価", "月販（下限）", "評価数", "★", "★1〜2の割合（%）", "梱包（推定）", "梱包重量",
                     "国内送料", "1688標準原価（元）", "テンプレ利益率", "実質利益率（広告5%）", "判定"],
             [12, 52, 9, 8, 8, 5, 8, 14, 8, 9, 9, 9, 10, 7], height=36)
    xs = sorted([x for x in EV if x["type"] == "dresser"], key=lambda x: -x["bought"])
    r = 5
    for x in xs:
        vals = [x["asin"], x["title"][:80], x["price"], x["bought"], x["reviews"], x["rating"], x.get("neg"),
                f"{x['L']:.0f}×{x['W']:.0f}×{x['H']:.0f}" if x.get("L") else "", x.get("kg"), x.get("dom"),
                x.get("cost"), x.get("tm"), x.get("r5"), x.get("judge")]
        fmts = [None, None, YEN, "#,##0", "#,##0", "0.0", "#,##0", None, KG, YEN, "#,##0.0", PCT, PCT, None]
        for j, (v, fm) in enumerate(zip(vals, fmts), start=1):
            cell = put(ws, f"{B.get_column_letter(j)}{r}", v, B.F_EST if j in (8, 9) else B.F_SMALL, fmt=fm,
                       align=B.LEFT_C if j == 2 else B.WRAP_C)
            if j == 1:
                cell.hyperlink = AMZ + x["asin"]
                cell.font = B.F_LINK
        r += 1
    r += 1
    put(ws, f"A{r}", "成立条件の試算（梱包と売価の組合せ・1688原価は仕様一致品の値。式は01シートと同じ）", B.F_SUB, border=False)
    r += 1
    cny = D.DRESSER["cny"]
    head_row(ws, r, ["梱包案", "長", "幅", "高", "重量", "売価", "国内送料", "テンプレ利益率", "実質利益率（広告5%）", "1個利益（広告5%）", "判定"], None)
    r += 1
    for name, (L, W, H, kg) in D.DRESSER["packs"]:
        for price in D.DRESSER["prices"]:
            e = A.econ(price, cny, L, W, H, kg, "F", va=D.DRESSER["va"])
            j = "GO" if e["tm"] >= 0.30 and e["r5"] >= 0.15 else ("COND" if e["tm"] >= 0.25 and e["r5"] >= 0.10 else "NO")
            for c_, v, fm in zip("ABCDEFGHIJK", (name, L, W, H, kg, price, e["dom"], e["tm"], e["r5"], e["p5"], j),
                                 (None, CM, CM, CM, KG, YEN, YEN, PCT, PCT, YEN, None)):
                put(ws, f"{c_}{r}", v, B.F_BOLD if c_ == "K" else B.F_SMALL, fmt=fm, align=B.WRAP_C)
            r += 1
    for txt, fill in (("GO", B.FILL_GO), ("COND", B.FILL_COND), ("NO", B.FILL_NO)):
        ws.conditional_formatting.add(f"K5:K{r}", CellIsRule(operator="equal", formula=[f'"{txt}"'], fill=fill))
        ws.conditional_formatting.add(f"N5:N{r}", CellIsRule(operator="equal", formula=[f'"{txt}"'], fill=fill))
    r += 1
    put(ws, f"A{r}", "1688の仕様一致品（LEDミラー付き・月销順）", B.F_SUB, border=False)
    r += 1
    head_row(ws, r, ["価格（元）", "月销", "出品タイトル（末尾）"], None)
    for o in sorted(FIN.get("dresser", []), key=lambda o: -o["sold"])[:15]:
        r += 1
        put(ws, f"A{r}", o["price"], B.F_SMALL, fmt="#,##0.0")
        put(ws, f"B{r}", o["sold"], B.F_SMALL, fmt="#,##0")
        ws.merge_cells(f"C{r}:N{r}")
        put(ws, f"C{r}", o["title"], B.F_SMALL, align=B.LEFT_C)
    r += 2
    put(ws, f"A{r}", "結論", B.F_SUB, border=False)
    for t in D.DRESSER["conclusion"]:
        r += 1
        ws.merge_cells(f"A{r}:N{r}")
        put(ws, f"A{r}", t, B.F_BASE, align=B.WRAP_TL, border=False)
        ws.row_dimensions[r].height = 32
    ws.freeze_panes = "C5"


# ---------------------------------------------------------------------------
def rival_dict(asin, f, cny):
    x, d = EVA.get(asin, {}), DET.get(asin, {})
    if not x:  # rival outside the screen (e.g. no purchase badge): page data and our own carton estimate
        L, W, H, kg = f["pkg"]
        x = dict(title=d.get("title", ""), price=d.get("price"), bought=d.get("bought") or 0, reviews=d.get("reviews"),
                 rating=d.get("rating"), L=L, W=W, H=H, kg=kg, pk_src="推定（自社企画の梱包）",
                 sub_cat=P_.sub_rank(d.get("bsr"))[0], sub_rank=P_.sub_rank(d.get("bsr"))[1])
    dims = A.parse_dims(d.get("dims")) or ["", "", ""]
    brand = (d.get("brand") or "").split(" ")[0] if d.get("brand") and "＆" not in d.get("brand", "") else ""
    return dict(name=x.get("title", d.get("title", ""))[:120], url=AMZ + asin, asin=asin,
                maker=d.get("seller") or brand or (d.get("maker") or "").split(" ")[0],
                price=x.get("price") or d.get("price"), units=x.get("bought"), reviews=x.get("reviews"), rating=x.get("rating"),
                pkg=(round(x.get("L", 0), 0), round(x.get("W", 0), 0), round(x.get("H", 0), 0)), pkg_kg=round(x.get("kg", 0), 1),
                prod=dims, prod_kg=d.get("weight", "")[:20], src=x.get("pk_src", ""), rank=f"{x.get('sub_cat') or ''} {x.get('sub_rank') or ''}位",
                neg=x.get("neg"), cny=cny)


def fill_rival(ws, b, rv, f):
    F12 = Font(name="Arial", size=12)
    put(ws, b["name"], rv["name"], B.F_BASE, align=B.LEFT_C, border=False)
    cell = put(ws, b["site"], rv["url"], B.F_URL, align=B.LEFT_C, border=False)
    cell.hyperlink = rv["url"]
    put(ws, b["maker"], rv["maker"], B.F_BASE, align=B.WRAP_C, border=False)
    put(ws, b["price"], f"¥{rv['price'] or 0:,.0f}（月{rv['units'] or 0}個以上購入・評価{rv['reviews'] or 0:,}件 ★{rv['rating']}）", B.F_BASE,
        align=B.WRAP_C, border=False)
    est = "推定" in rv["src"]
    for ref, v in zip(b["pkg"], rv["pkg"]):
        put(ws, ref, v, B.F_EST if est else B.F_BASE, fmt=CM, align=B.WRAP_C, border=False)
    put(ws, b["pkg_kg"], rv["pkg_kg"], B.F_EST if est else B.F_BASE, fmt=KG, align=B.WRAP_C, border=False)
    for ref, v in zip(b["prod"], rv["prod"]):
        put(ws, ref, v, B.F_SMALL, fmt=CM if not isinstance(v, str) else None, align=B.WRAP_C, border=False)
    put(ws, b["prod_kg"], rv["prod_kg"], B.F_SMALL, align=B.WRAP_C, border=False)
    put(ws, b["wcheck"], f"=IF({b['pkg_kg']}<=50,\"OK\",\"50kg超→分割\")", B.F_SMALL, align=B.WRAP_C, border=False)
    note = (f"【実データ 2026-10-21】売れ筋 {rv['rank']}／★1〜2の評価 {rv['neg'] if rv['neg'] is not None else '-'}%。"
            f"梱包：{rv['src']}（青字＝推定・要実測）。原価：1688の仕様一致品の中央値。")
    put(ws, b["note"], note, B.F_NOTE, align=B.WRAP_TL, border=False)
    put(ws, b["asin"], rv["asin"], B.F_SMALL, align=B.WRAP_C, border=False)
    r = b["vrow"]
    put(ws, f"G{r}", rv["price"], B.F_BASE, fmt='"¥"#,##0_);[Red]\\("¥"#,##0\\)', border=False)
    put(ws, f"H{r}", round(rv["cny"], 1), B.F_EST, fmt=CNY, border=False)
    put(ws, f"I{r}", f"={P['fx']}", F12, align=B.WRAP_C, border=False)
    L, W, H = b["pkg"]
    put(ws, f"K{r}", ship_f(L, W, H, b["pkg_kg"], f"IF({b['pkg_kg']}>50,ROUNDUP({b['pkg_kg']}/50,0),IF({b['sum_ref']}>260,2,1))"),
        F12, fmt='"¥"#,##0', border=False)
    put(ws, f"O{r}", f"={P['fr_t']}", F12, fmt="#,##0", border=False)
    put(ws, f"T{r}", rv["units"], F12, align=Alignment(horizontal="right"), border=False)
    put(ws, f"V{r}", rv["rank"].strip(), B.F_SMALL, align=B.WRAP_C, border=False)
    put(ws, f"W{r}", f["order"], F12, align=Alignment(horizontal="right"), border=False)


def add_plan_block(ws, f, ov):
    s = 77
    blue = B.PatternFill("solid", fgColor="BDD7EE")
    merge_put(ws, f"A{s}:O{s}", "【Claude追記】自社企画品（差別化仕様）の収益試算 ― 左：企画仕様、右上：テンプレート式、右下：実質（天瞳の実料金）",
              B.F_BOLD, fill=blue, align=B.LEFT_C)
    left = [
        (s + 1, "企画品名", f"={OV}!B{ov}", None), (s + 2, "想定販売価格（定常）", f"={OV}!E{ov}", YEN),
        (s + 3, "導入価格（レビュー獲得期）", f["launch"], YEN), (s + 4, "1688仕入原価（元）", f"={OV}!F{ov}", CNY),
        (s + 5, "レート", f"={P['fx']}", "0.00"), (s + 6, "円の商品原価", f"=E{s + 4}*E{s + 5}", YEN),
        (s + 7, "付加価値費（説明書・補強部品等）", f"={OV}!G{ov}", YEN), (s + 8, "梱包 長／幅／高（cm）", None, None),
        (s + 9, "梱包重量（kg）・個口", f"={OV}!K{ov}", KG), (s + 10, "CBM（体積）", f"={OV}!M{ov}", M3),
        (s + 11, "三辺合計（cm）", f"={OV}!N{ov}", CM), (s + 12, "国内送料（天瞳・佐川）", f"={OV}!O{ov}", YEN),
        (s + 13, "月販目標（3年目）", f"={OV}!AB{ov}", "#,##0"), (s + 14, "初回発注数", f["order"], "#,##0"),
        (s + 15, "初回仕入資金（原価・諸費用・付加価値・国際送料17,325円/m³）", None, YEN),
    ]
    for r, label, v, fmt in left:
        ws.merge_cells(f"A{r}:D{r}")
        put(ws, f"A{r}", label, B.F_SMALL, fill=B.FILL_SUB, align=B.LEFT_C)
        if r == s + 1:
            ws.merge_cells(f"E{r}:O{r}")
            put(ws, f"E{r}", v, B.F_BOLD, align=B.LEFT_C)
        elif r == s + 8:
            for c_, oc in zip("EFG", "HIJ"):
                put(ws, f"{c_}{r}", f"={OV}!{oc}{ov}", fmt=CM)
        elif r == s + 9:
            put(ws, f"E{r}", v, fmt=fmt)
            put(ws, f"F{r}", f"={OV}!L{ov}", fmt='0"個口"')
        elif r == s + 15:
            put(ws, f"E{r}", f"=E{s + 14}*(E{s + 6}*(1+{P['cn']})+E{s + 7}+{P['fr_t']}*E{s + 10})", B.F_BOLD, fmt=YEN)
        else:
            put(ws, f"E{r}", v, fmt=fmt, fill=B.FILL_IN if r in (s + 3, s + 14) else None)
    merge_put(ws, f"I{s + 2}:O{s + 2}", "テンプレート式（天瞳リサーチ表と同じ計算）", B.F_BOLD, fill=B.FILL_SUB, align=B.LEFT_C)
    tm = [
        (s + 3, "販売価額", f"=E{s + 2}", YEN), (s + 4, "円の商品原価", f"=E{s + 6}", YEN), (s + 5, "国内送料", f"=E{s + 12}", YEN),
        (s + 6, "Amazon手数料（16.5%）", f"=M{s + 3}*{P['tfee']}", YEN), (s + 7, "消費税（原価×10%）", f"=M{s + 4}*{P['tvat']}", YEN),
        (s + 8, "関税（原価×10%）", f"=M{s + 4}*{P['tduty']}", YEN), (s + 9, "国際運送（17,325円×CBM）", f"={P['fr_t']}*E{s + 10}", YEN),
        (s + 10, "粗利（テンプレ）", f"=M{s + 3}-M{s + 4}-M{s + 5}-M{s + 6}-M{s + 7}-M{s + 8}-M{s + 9}", YEN),
        (s + 11, "利益率（テンプレ）", f"=M{s + 10}/M{s + 3}", PCT),
    ]
    rl = [
        (s + 13, "売上（税抜）", f"={OV}!R{ov}", YEN), (s + 14, "販売手数料（カテゴリ別）", f"={OV}!S{ov}", YEN),
        (s + 15, "国内送料", f"=E{s + 12}", YEN), (s + 16, "返品・破損（売上×4%）", f"=M{s + 13}*{P['ret']}", YEN),
        (s + 17, "保管料（CBM×77円×45日）", f"={OV}!T{ov}", YEN), (s + 18, "仕入原価（中国側諸費用8%込）", f"=E{s + 6}*(1+{P['cn']})", YEN),
        (s + 19, "付加価値費", f"=E{s + 7}", YEN), (s + 20, "国際送料（13,125円/m³）", f"={P['fr']}*E{s + 10}", YEN),
        (s + 21, "広告費前利益", f"=M{s + 13}-SUM(M{s + 14}:M{s + 20})", YEN),
        (s + 22, "実質利益（広告10%：立ち上げ期）", f"=M{s + 21}-M{s + 13}*{P['ad10']}", YEN),
        (s + 23, "実質利益（広告5%：1位定着後）", f"=M{s + 21}-M{s + 13}*{P['ad5']}", YEN),
        (s + 24, "判定（テンプレ≥30%かつ実質≥15%→GO）", None, None),
        (s + 25, "月利（3年目の月販×実質利益・広告5%）", f"=E{s + 13}*M{s + 23}", YEN),
    ]
    for r, label, v, fmt in tm + rl:
        ws.merge_cells(f"I{r}:L{r}")
        put(ws, f"I{r}", label, B.F_SMALL, fill=B.FILL_SUB, align=B.LEFT_C)
        if r == s + 24:
            put(ws, f"M{r}", f"=IF(AND(M{s + 11}>={P['go_t']},N{s + 23}>={P['go_r']}),\"GO\",IF(AND(M{s + 11}>={P['c_t']},N{s + 23}>={P['c_r']}),\"COND\",\"NO\"))",
                B.F_RED, align=B.WRAP_C)
        else:
            put(ws, f"M{r}", v, B.F_BOLD if r in (s + 10, s + 11, s + 23, s + 25) else B.F_BASE, fmt=fmt)
    merge_put(ws, f"I{s + 12}:O{s + 12}", "実質（天瞳の実料金：佐川送料・保管77円/m³日・海上13,125円/m³）", B.F_BOLD, fill=B.FILL_SUB, align=B.LEFT_C)
    for r in (s + 22, s + 23):
        put(ws, f"N{r}", f"=M{r}/M{s + 13}", B.F_BOLD, fmt=PCT)
        put(ws, f"O{r}", "利益率", B.F_NOTE, border=False)
    hx = "".join("%02X" % b for b in f["kw_cn"].encode("gbk"))
    for r, label, text, url in (
            (s + 17, "1688で同等品を検索（クリック）", f"▶ 1688：{f['kw_cn']}", f"https://www.1688.com/market/-{hx}.html"),
            (s + 18, "Amazonで競合を検索（クリック）", f"▶ Amazon：{f['kw_jp']}", "https://www.amazon.co.jp/s?k=" + quote(f["kw_jp"]))):
        ws.merge_cells(f"A{r}:D{r}")
        put(ws, f"A{r}", label, B.F_SMALL, fill=B.FILL_SUB, align=B.LEFT_C)
        ws.merge_cells(f"E{r}:H{r}")
        cell = put(ws, f"E{r}", text, B.F_LINK, align=B.LEFT_C)
        cell.hyperlink = url
    # 1688 spec-matched offers used for the cost
    r0 = s + 28
    merge_put(ws, f"A{r0}:O{r0}", "【原価の根拠】1688の仕様一致品（月销の多い順・上位10件）", B.F_BOLD, fill=blue, align=B.LEFT_C)
    for i, o in enumerate(evidence(f)):
        r = r0 + 1 + i
        put(ws, f"A{r}", o["price"], B.F_SMALL, fmt=CNY)
        put(ws, f"B{r}", o["sold"], B.F_SMALL, fmt='"月销"#,##0')
        ws.merge_cells(f"C{r}:O{r}")
        put(ws, f"C{r}", o["title"], B.F_SMALL, align=B.LEFT_C)


def evidence(f, n=10):
    """1688 offers behind a finalist's cost: spec-matched searches first, then same-type offers whose
    title matches the spec's first pattern; most sold first."""
    rx = re.compile(T2[f["fin_key"]][1][0])
    seen, out = set(), []
    for o in FIN.get(f["fin_key"], []) + [o for o in CN.get(f["type"], []) if rx.search(o["title"])]:
        key = (o["price"], o["sold"])  # the same offer appears under several shop names
        if key not in seen:
            seen.add(key)
            out.append(o)
    return sorted(out, key=lambda o: -o["sold"])[:n]


def review_texts(f):
    """Complaint/praise counts for the rivals of this type (all leaders' product pages)."""
    revs = []
    for a, x in DET.items():
        if a in PRODS and KW2T.get(PRODS[a]["kws"][0]) == f["type"]:
            revs += [tuple(v) for v in x.get("reviews_txt") or []]
    neg, pos, ex = C.mine(revs)
    return len(revs), neg, pos, ex


def build_product_sheet(wb, tpl, f, ov):
    ws = wb.copy_worksheet(tpl)
    ws.title = f["sheet"]
    B.copy_conditional_formatting(tpl, ws)
    ws["A1"].value = f"リサーチ｜{f['code']} {f['title']}"
    ws["A1"].font = Font(name=B.FONT, size=20, bold=True)
    for b, asin, cny in zip(B.BLOCKS, f["rivals"], f["rival_cny"]):
        fill_rival(ws, b, rival_dict(asin, f, cny), f)
    ws["E9"].value = ("※ランキング30/60/90日平均はKeepa等で取得してください（緑のまま）。V列は2026/10/21時点のサブカテゴリ順位。"
                      "販売価格・購入数・評価はAmazon実データ。梱包寸法と原価は推定（青字）。")
    ws["E9"].font = B.F_NOTE
    n, neg, pos, ex = review_texts(f)
    auto_m = "【商品ページのレビュー集計（★4〜5）】" + "、".join(f"{k}{v}件" for k, v in pos.most_common(4)) + f"（対象{n}件）\n"
    auto_d = "【★1〜3レビューの不満の集計】" + "、".join(f"{k}{v}件" for k, v in neg.most_common(5)) + "\n"
    for ref, txt in (("B31", auto_m + f["merits"]), ("B36", auto_d + f["demerits"]), ("B41", f["ours"]), ("B47", f["latent"])):
        c = ws[ref]
        c.value = txt
        c.font = Font(name=B.FONT, size=10)
        c.alignment = B.WRAP_TL
    for r in list(range(31, 35)) + list(range(36, 40)):
        ws.row_dimensions[r].height = 30
    ws["Y1"].value = f["memo"]
    ws["Y1"].font = Font(name=B.FONT, size=10)
    ws["Y1"].alignment = B.WRAP_TL
    for col in ("Y", "Z", "AA", "AB"):
        ws.column_dimensions[col].width = 16
    add_plan_block(ws, f, ov)
    return ws


# ---------------------------------------------------------------------------
def build_strategy(ws):
    ws.title = "05_1位獲得戦略"
    put(ws, "A1", "ランキング1位を取るための戦略（最終候補・Amazon実データ）", B.F_TITLE, border=False)
    head_row(ws, 4, ["詳細シート", "役割", "狙うキーワード・ノード", "現在の1位（価格・月販）", "1位を取れる根拠", "差別化仕様（必須）",
                     "価格戦略（導入→定常）", "1位ライン（月販）", "立ち上げ施策", "主なリスクと対策"],
             [16, 10, 30, 36, 40, 40, 20, 10, 34, 34], height=36)
    for i, f in enumerate(D.FINALISTS):
        r = 5 + i
        s = f["strategy"]
        vals = [f["sheet"], f["role"], s["kw"], s["leader"], s["why"], s["spec"], s["price"], f["top_units"], s["launch"], s["risk"]]
        for j, v in enumerate(vals, start=1):
            cell = put(ws, f"{B.get_column_letter(j)}{r}", v, B.F_SMALL, align=B.WRAP_TL, fmt="#,##0" if j == 8 else None)
            if j == 1:
                cell.font = B.F_LINK
                link_sheet(cell, f"'{f['sheet']}'!A1")
        ws.row_dimensions[r].height = 110
    ws.freeze_panes = "B5"


def build_roadmap(ws, rowmap):
    ws.title = "06_年商5億ロードマップ"
    put(ws, "A1", "年商5億円・年収5,000万円への到達計画（Amazon＋楽天・Yahoo!）", B.F_TITLE, border=False)
    put(ws, "A2", "月販は計画値（黄色＝入力、各市場の1位の月販以下）。1年目は広告10%、2年目以降は広告5%。固定費・法人税は含まない。", B.F_NOTE, border=False)
    head_row(ws, 4, ["商品", "役割", "詳細シート", "定常売価", "1個利益（広告10%）", "1個利益（広告5%）", "1個仕入総額",
                     "1年目 月販", "2年目 月販", "3年目 月販", "1年目 年商", "2年目 年商", "3年目 年商", "1年目 年利", "2年目 年利", "3年目 年利"],
             [40, 10, 16, 10, 10, 10, 10, 8, 8, 8, 14, 14, 14, 13, 13, 13], height=36)
    r = 5
    for f in D.FINALISTS:
        o = rowmap[f["id"]]
        put(ws, f"A{r}", f"={OV}!B{o}", B.F_SMALL, align=B.LEFT_C)
        put(ws, f"B{r}", f["role"], B.F_SMALL, align=B.WRAP_C)
        cell = put(ws, f"C{r}", f["sheet"], B.F_LINK)
        link_sheet(cell, f"'{f['sheet']}'!A1")
        put(ws, f"D{r}", f"={OV}!E{o}", fmt=YEN)
        put(ws, f"E{r}", f"={OV}!V{o}", fmt=YEN)
        put(ws, f"F{r}", f"={OV}!W{o}", fmt=YEN)
        put(ws, f"G{r}", f"={OV}!U{o}", fmt=YEN)
        for col, v in zip("HIJ", f["y"]):
            put(ws, f"{col}{r}", v, fmt="#,##0", fill=B.FILL_IN)
        for col, uc in zip("KLM", "HIJ"):
            put(ws, f"{col}{r}", f"=D{r}*{uc}{r}*12", fmt=YEN)
        put(ws, f"N{r}", f"=E{r}*H{r}*12", fmt=YEN)
        put(ws, f"O{r}", f"=F{r}*I{r}*12", fmt=YEN)
        put(ws, f"P{r}", f"=F{r}*J{r}*12", fmt=YEN)
        r += 1
    last, t = r - 1, r
    put(ws, f"A{t}", "Amazon 合計", B.F_BOLD, fill=B.FILL_SUB)
    for col in "HIJKLMNOP":
        put(ws, f"{col}{t}", f"=SUM({col}5:{col}{last})", B.F_BOLD, fmt=YEN if col in "KLMNOP" else "#,##0", fill=B.FILL_SUB)
    u = t + 2
    put(ws, f"A{u}", "楽天・Yahoo!の上乗せ率（Amazon比）", B.F_BOLD)
    for col, v in zip("KLM", (0, 0.15, 0.25)):
        put(ws, f"{col}{u}", v, fmt=PCT, fill=B.FILL_IN)
    put(ws, f"A{u + 1}", "総年商（全モール）", B.F_BOLD, fill=B.FILL_KEY)
    put(ws, f"A{u + 2}", "総年利（営業利益・固定費前）", B.F_BOLD, fill=B.FILL_KEY)
    for col, pc in zip("KLM", "NOP"):
        put(ws, f"{col}{u + 1}", f"={col}{t}*(1+{col}{u})", B.F_BOLD, fmt=YEN, fill=B.FILL_KEY)
        put(ws, f"{pc}{u + 2}", f"={pc}{t}*(1+{col}{u})", B.F_BOLD, fmt=YEN, fill=B.FILL_KEY)
    put(ws, f"A{u + 4}", "目標年商", B.F_BOLD)
    put(ws, f"K{u + 4}", 500000000, fmt=YEN, fill=B.FILL_IN)
    put(ws, f"A{u + 5}", "3年目の達成率（総年商÷目標）", B.F_BOLD)
    put(ws, f"K{u + 5}", f"=M{u + 1}/K{u + 4}", B.F_RED, fmt=PCT)
    put(ws, f"A{u + 6}", "3年目に必要な在庫資金（月販×3か月×1個仕入総額）", B.F_BOLD)
    put(ws, f"K{u + 6}", f"=SUMPRODUCT(G5:G{last},J5:J{last})*3", B.F_BOLD, fmt=YEN)
    put(ws, f"A{u + 7}", "3年目の年間輸入体積（m³）→ 40ftコンテナ本数", B.F_BOLD)
    put(ws, f"K{u + 7}", "=" + "+".join(f"{OV}!M{rowmap[f['id']]}*J{5 + i}*12" for i, f in enumerate(D.FINALISTS)), fmt='#,##0"m³"')
    put(ws, f"L{u + 7}", f"=K{u + 7}/67", fmt='0"本"')
    for i, tx in enumerate(D.ROADMAP_NOTES):
        rr = u + 9 + i
        ws.merge_cells(f"A{rr}:P{rr}")
        put(ws, f"A{rr}", tx, B.F_SMALL, align=B.WRAP_TL, border=False)
        ws.row_dimensions[rr].height = 30
    ws.freeze_panes = "B5"
    return dict(total=t, uplift=u)


def build_summary(ws, rowmap, rm):
    ws.title = "00_結論"
    ws.column_dimensions["A"].width = 44
    for col, w in zip("BCDEFGHIJK", (12, 10, 11, 11, 11, 11, 11, 12, 10, 20)):
        ws.column_dimensions[col].width = w
    put(ws, "A1", "大型OEM 商品化判断 v4（Amazon・1688実データ × 天瞳の実料金）", B.F_TITLE, border=False)
    put(ws, "A2", D.SUMMARY_SUB, B.F_NOTE, border=False)
    put(ws, "A4", "結論（要点）", B.F_SUB, border=False)
    for i, t in enumerate(D.SUMMARY):
        r = 5 + i
        ws.merge_cells(f"A{r}:K{r}")
        put(ws, f"A{r}", t, B.F_BASE, align=B.WRAP_TL, border=False)
        ws.row_dimensions[r].height = 46
    r1 = 6 + len(D.SUMMARY)
    put(ws, f"A{r1}", "最終候補（数値は01シートから自動参照）", B.F_SUB, border=False)
    hdr = ["商品", "役割", "判定", "定常売価", "テンプレ利益率", "実質利益率（広告5%）", "1個利益（広告5%）", "価格差", "競合1位の月販",
           "月利（3年目）", "詳細シート"]
    for j, h in enumerate(hdr, start=1):
        put(ws, f"{B.get_column_letter(j)}{r1 + 1}", h, B.F_HEAD, fill=B.FILL_HEAD, align=B.WRAP_C)
    for i, f in enumerate(D.FINALISTS):
        r, o = r1 + 2 + i, rowmap[f["id"]]
        put(ws, f"A{r}", f"={OV}!B{o}", B.F_SMALL, align=B.LEFT_C)
        put(ws, f"B{r}", f["role"], B.F_SMALL, align=B.WRAP_C)
        put(ws, f"C{r}", f"={OV}!Y{o}", B.F_BOLD, align=B.WRAP_C)
        put(ws, f"D{r}", f"={OV}!E{o}", fmt=YEN)
        put(ws, f"E{r}", f"={OV}!Q{o}", fmt=PCT)
        put(ws, f"F{r}", f"={OV}!X{o}", B.F_BOLD, fmt=PCT)
        put(ws, f"G{r}", f"={OV}!W{o}", fmt=YEN)
        put(ws, f"H{r}", f"={OV}!Z{o}", fmt='0.0"倍"')
        put(ws, f"I{r}", f"={OV}!AA{o}", fmt="#,##0")
        put(ws, f"J{r}", f"={OV}!AC{o}", B.F_BOLD, fmt=YEN)
        cell = put(ws, f"K{r}", f["sheet"], B.F_LINK)
        link_sheet(cell, f"'{f['sheet']}'!A1")
        ws.row_dimensions[r].height = 28
    last = r1 + 1 + len(D.FINALISTS)
    for txt, fill in (("GO", B.FILL_GO), ("COND", B.FILL_COND), ("NO", B.FILL_NO)):
        ws.conditional_formatting.add(f"C{r1 + 2}:C{last}", CellIsRule(operator="equal", formula=[f'"{txt}"'], fill=fill))
    r2 = last + 2
    put(ws, f"A{r2}", "年商5億の到達見込み（06シートから自動参照）", B.F_SUB, border=False)
    for j, h in enumerate(["", "1年目", "2年目", "3年目"], start=1):
        put(ws, f"{B.get_column_letter(j)}{r2 + 1}", h, B.F_HEAD, fill=B.FILL_HEAD, align=B.WRAP_C)
    u, RM = rm["uplift"], "'06_年商5億ロードマップ'"
    put(ws, f"A{r2 + 2}", "総年商（Amazon＋楽天・Yahoo!）", B.F_BOLD)
    put(ws, f"A{r2 + 3}", "総年利（営業利益・固定費前）", B.F_BOLD)
    for col, s_, p_ in (("B", "K", "N"), ("C", "L", "O"), ("D", "M", "P")):
        put(ws, f"{col}{r2 + 2}", f"={RM}!{s_}{u + 1}", B.F_BOLD, fmt=YEN)
        put(ws, f"{col}{r2 + 3}", f"={RM}!{p_}{u + 2}", B.F_BOLD, fmt=YEN)
    put(ws, f"A{r2 + 4}", "3年目の達成率（年商5億円比）", B.F_BOLD)
    put(ws, f"D{r2 + 4}", f"={RM}!K{u + 5}", B.F_RED, fmt=PCT)
    r3 = r2 + 6
    put(ws, f"A{r3}", "次にやること（優先順）", B.F_SUB, border=False)
    for i, t in enumerate(D.TODO):
        r = r3 + 1 + i
        ws.merge_cells(f"A{r}:K{r}")
        put(ws, f"A{r}", t, B.F_BASE, align=B.WRAP_TL, border=False)
        ws.row_dimensions[r].height = 30


def build_raw_1688(ws):
    ws.title = "生データ_1688"
    head_row(ws, 1, ["区分", "タイプ／候補", "1688検索語", "出品（タイトル末尾）", "価格（元）", "月销"], [12, 14, 26, 70, 10, 10])
    tj = {v["id"]: v["jp"] for v in D.TYPES}
    for src, data in (("最終候補・仕様一致", FIN), ("タイプ別検索", CN)):
        for k, offers in data.items():
            for o in sorted(offers, key=lambda o: -o["sold"]):
                ws.append([src, tj.get(k, k), o.get("kw", ""), o["title"], o["price"], o["sold"]])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:F{ws.max_row}"


def build_sources(ws):
    ws.title = "出典"
    put(ws, "A1", "出典・取得方法", B.F_TITLE, border=False)
    head_row(ws, 3, ["No", "区分", "内容", "URL"], [5, 14, 90, 60])
    for i, (cat, text, url) in enumerate(D.SOURCES, start=1):
        r = 3 + i
        put(ws, f"A{r}", i)
        put(ws, f"B{r}", cat, B.F_SMALL)
        put(ws, f"C{r}", text, B.F_SMALL, align=B.WRAP_TL)
        cell = put(ws, f"D{r}", url, B.F_URL, align=B.WRAP_TL)
        if url.startswith("http"):
            cell.hyperlink = url


def main(template, out, extra=None):
    """extra(wb) may add sheets; they are placed after the summary sheets."""
    wb = load_workbook(template)
    tpl = wb["リサーチ表"]
    names = ["00_結論", "01_最終候補", "02_全商品スクリーニング", "03_タイプ別集計", "04_ドレッサー検証",
             "05_1位獲得戦略", "06_年商5億ロードマップ", "前提条件", "生データ_1688", "出典"]
    S = {n: wb.create_sheet(n) for n in names}
    build_assumptions(S["前提条件"])
    rowmap, _ = build_overview(S["01_最終候補"])
    build_screen(S["02_全商品スクリーニング"])
    build_types(S["03_タイプ別集計"])
    build_dresser(S["04_ドレッサー検証"])
    build_strategy(S["05_1位獲得戦略"])
    rm = build_roadmap(S["06_年商5億ロードマップ"], rowmap)
    build_summary(S["00_結論"], rowmap, rm)
    build_raw_1688(S["生データ_1688"])
    build_sources(S["出典"])
    sheets = [build_product_sheet(wb, tpl, f, rowmap[f["id"]]) for f in D.FINALISTS]
    tpl.title = "テンプレート(原本)"
    head = [S[n] for n in names[:7]]
    added = extra(wb) if extra else []
    wb._sheets = head + added + sheets + [S["前提条件"], S["生データ_1688"], S["出典"], tpl]
    wb.active = 0
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = ws is S["00_結論"]
    wb.save(out)
    for f in D.FINALISTS + D.REJECTED:
        L, W, H, kg = f["pkg"]
        e = A.econ(f["price"], f["cny"], L, W, H, kg, f["fee"], va=f["va"], n=f["parcels"])
        print(f"{f['id']} tm {e['tm']*100:5.1f}% r5 {e['r5']*100:5.1f}% p5 {e['p5']:7.0f} dom {e['dom']} {f['name'][:30]}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
