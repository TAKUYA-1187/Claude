# -*- coding: utf-8 -*-
"""Build the v3 workbook (real Amazon/1688 data) on the user's 天瞳 template.

usage: OEM_DATA=<dir with scan.json/detail.json/detail_extra.json> \
       python3 build_v3.py <template.xlsx> <out.xlsx>
Reuses the v2 builders (template sheets, assumptions, parcels, overview, strategy).
"""
import os
import sys
from urllib.parse import quote

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "v2"))
import analyze as A  # noqa: E402
import build_v2 as B  # noqa: E402
import data_v3 as D  # noqa: E402
from keywords import KEYWORDS  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from openpyxl.formatting.rule import FormulaRule  # noqa: E402

B.D = D
B.FIN_IDX = {c["id"]: i for i, c in enumerate(D.FINALISTS)}
B.ORDER.update({"採用（HERO・梱包0.10m³以内が条件）": 0, "採用（梱包170サイズ化が条件）": 2})
put, head_row = B.put, B.head_row
YEN, PCT = B.YEN, B.PCT


def build_market(ws):
    ws.title = "01_市場スキャン121"
    put(ws, "A1", "Amazon.co.jp 市場スキャン（大型・中型121キーワード／2026-10-20実測）", B.F_TITLE, border=False)
    put(ws, "A2", "配送先を東京都に設定し、各キーワードの検索結果2ページ（約100件）を取得。月販＝『過去1か月でN点以上購入』バッジの下限値の合計（50個未満の商品は0扱い＝控えめな値）。"
                  "タイトルがキーワードに合わない商品・¥2,500未満の付属品は除外。1688原価＝市場価格の6〜50%に入る1688出品の中央値。", B.F_NOTE, border=False)
    head_row(ws, 4, ["区分", "キーワード", "対象出品数", "バッジ付き", "月販合計（下限）", "月商合計（下限）", "中央価格",
                     "1位の月販", "1位の価格", "1位の評価数", "1位の★", "1位の商品", "2位の商品（月販・価格）", "上位10の平均★",
                     "上位10の国内大手数", "1688原価中央値（元）", "1688対象出品数", "1688月销合計", "Amazon検索", "1688検索"],
             [10, 22, 7, 7, 9, 12, 9, 7, 9, 8, 6, 46, 40, 7, 7, 9, 7, 9, 10, 10], height=44)
    rows = sorted((A.market(kw, D.SCAN[kw]) for _, kw, _, _ in KEYWORDS if kw in D.SCAN),
                  key=lambda m: -m["market_yen"])
    for i, m in enumerate(rows):
        r = 5 + i
        t = m["top1"] or {}
        t2 = m["top10"][1] if len(m["top10"]) > 1 else {}
        vals = [m["group"], m["kw"], m["n"], m["n_badged"], m["market_u"], m["market_yen"], m["price_med"],
                t.get("bought"), t.get("price"), t.get("reviews"), t.get("rating"), t.get("title", "")[:80],
                f"{t2.get('bought', 0)}+/月 ¥{t2.get('price') or 0:,.0f} {t2.get('title', '')[:40]}" if t2 else "",
                round(m["top10_rating"], 2) if m["top10_rating"] else None, m["bigjp_top10"],
                round(m["cost"], 1) if m["cost"] else None, m["n_offers"], m["offers_sold"]]
        fmts = [None, None, "#,##0", "#,##0", "#,##0", YEN, YEN, "#,##0", YEN, "#,##0", "0.0", None, None, "0.00",
                "0", "#,##0.0", "#,##0", "#,##0"]
        for j, (v, f) in enumerate(zip(vals, fmts), start=1):
            cell = put(ws, f"{B.get_column_letter(j)}{r}", v, B.F_SMALL, fmt=f,
                       align=B.WRAP_TL if j in (12, 13) else None)
            if j == 12 and t.get("asin"):
                cell.hyperlink = D.AMZ + t["asin"]
                cell.font = B.F_LINK
        c = put(ws, f"S{r}", "Amazon", B.F_LINK)
        c.hyperlink = "https://www.amazon.co.jp/s?k=" + quote(m["kw"])
        c = put(ws, f"T{r}", "1688", B.F_LINK)
        c.hyperlink = "https://www.1688.com/market/-" + "".join("%02X" % b for b in m["cn"].encode("gbk")) + ".html"
        ws.row_dimensions[r].height = 30
    last = 4 + len(rows)
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:T{last}"
    ws.conditional_formatting.add(f"F5:F{last}", FormulaRule(formula=["F5>=30000000"], fill=B.FILL_GO))
    return rows


def build_raw_amazon(ws):
    ws.title = "生データ_Amazon"
    head_row(ws, 1, ["キーワード", "ASIN", "商品名", "価格", "★", "評価数", "過去1か月の購入（下限）", "関連性"],
             [20, 12, 80, 9, 5, 8, 10, 8])
    r = 2
    for _, kw, _, _ in KEYWORDS:
        rec = D.SCAN.get(kw)
        if not rec:
            continue
        rx = A.re.compile(A.REL.get(kw, A.re.escape(kw.split()[0])))
        for it in sorted(rec["items"], key=lambda x: -x["bought"]):
            ok = bool(it.get("price")) and it["price"] >= A.MIN_PRICE and bool(rx.search(it["title"]))
            ws.append([kw, it["asin"], it["title"][:120], it.get("price"), it.get("rating"), it.get("reviews"),
                       it["bought"], "対象" if ok else "除外"])
            ws.cell(r, 2).hyperlink = D.AMZ + it["asin"]
            r += 1
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:H{r - 1}"


def build_raw_1688(ws):
    ws.title = "生データ_1688"
    head_row(ws, 1, ["Amazonキーワード", "1688キーワード", "出品（タイトル末尾）", "価格（元）", "月销"], [20, 16, 60, 10, 10])
    for _, kw, cn, _ in KEYWORDS:
        rec = D.SCAN.get(kw)
        if not rec:
            continue
        for o in rec.get("offers", []):
            ws.append([kw, cn, o["title"], o["price"], o["sold"]])
    ws.freeze_panes = "A2"


def build_summary(ws, rowmap, last_ov, rm, markets):
    ws.title = "00_結論"
    ws.column_dimensions["A"].width = 46
    for col, w in zip("BCDEFGHIJK", (14, 14, 11, 11, 11, 11, 11, 12, 9, 20)):
        ws.column_dimensions[col].width = w
    put(ws, "A1", "大型OEM 商品化判断 v3（Amazon・1688の実データ／天瞳リサーチ表形式）", B.F_TITLE, border=False)
    put(ws, "A2", "2026-10-20 取得：Amazon.co.jp 121キーワード×2ページ（配送先 東京都）＋上位商品ページ約200件、1688 市場ページ121件。月販は購入バッジの下限値。",
        B.F_NOTE, border=False)
    put(ws, "A4", "結論（要点）", B.F_SUB, border=False)
    big = [m for m in markets if m["market_yen"] >= 30_000_000]
    pts = [
        f"1. 前回までは検索スニペットからの推定でしたが、今回はAmazonの実データです。121キーワード中、月商3,000万円以上（下限）の市場は{len(big)}個。"
        "大型品の1位の月販は多くが月200〜1,000個で、『1商品で月数千個』は現実的ではありません。",
        "2. 天瞳の国際送料（17,325円/m³）と国内送料では、¥1万未満・かさばる商品はほぼ赤字です（キャットタワー¥7千帯、収納ボックス、チェア、ペットカート、ゴルフマット等）。"
        "勝てるのは『単価¥1.2万以上で体積0.10m³前後以下』の市場です。",
        "3. 採用はHERO3品（ツイストステッパー・ベビーサークル・宅配ボックス ポスト一体型）＋主力2品（キャンプコット・猫ケージ3段）＋条件付き1品（フィットネスバイク）＋サポート1品（突っ張りタワー）。",
        "4. いずれも上位品の低評価レビューに具体的な不満（耐久・音・扉の単品販売・吸盤・説明書・組立・ロック・検品）があり、仕様で解消できます。"
        "ステッパーは評価106件の新規品が月1,000個以上売れており、新規参入で上位を取れる実例があります。",
        "5. 前回HEROにした大型猫キャットタワーと猫トイレ収納は、実データでは売れ筋が¥7千帯または月販ほぼゼロのため外しました。宅配ボックスはポスト一体型に切り替えました。",
        "6. 年商5億円は、7品をそれぞれ1位ライン近くまで伸ばし、楽天・Yahoo!を加えて3年目に届く計画です（04シート）。1品でも1位ラインに届かなければ未達になるため、"
        "1年目は上位3品に資金を集中し、Keepaで月販を確認しながら残りを投入する段階戦略を推奨します。",
        "7. 未確定：1688は市場ページの価格で、正式見積ではありません。梱包寸法はタンスのゲンの実寸1件以外は推定（青字）です。天瞳の国内送料表も未入手（前提条件に仮置き）。",
    ]
    for i, t in enumerate(pts):
        r = 5 + i
        ws.merge_cells(f"A{r}:K{r}")
        put(ws, f"A{r}", t, B.F_BASE, align=B.WRAP_TL, border=False)
        ws.row_dimensions[r].height = 48
    r1 = 13
    put(ws, f"A{r1}", "採用ラインナップ（数値は02シートから自動参照）", B.F_SUB, border=False)
    hdr = ["商品", "役割", "最終判定", "定常売価", "テンプレ利益率", "実質利益率（広告5%）", "1個利益（広告5%）",
           "月販目標（3年目）", "月利目標（3年目）", "1位可能性", "詳細シート"]
    for j, h in enumerate(hdr, start=1):
        put(ws, f"{B.get_column_letter(j)}{r1 + 1}", h, B.F_HEAD, fill=B.FILL_HEAD, align=B.WRAP_C)
    OV = B.OV
    for i, c in enumerate(D.FINALISTS):
        r = r1 + 2 + i
        o = rowmap[c["id"]]
        put(ws, f"A{r}", f"={OV}!B{o}", B.F_SMALL, align=B.LEFT_C)
        put(ws, f"B{r}", c["role"], B.F_SMALL, align=B.WRAP_C)
        put(ws, f"C{r}", f"={OV}!W{o}", B.F_SMALL, align=B.WRAP_C)
        put(ws, f"D{r}", f"={OV}!E{o}", fmt=YEN)
        put(ws, f"E{r}", f"={OV}!N{o}", fmt=PCT)
        put(ws, f"F{r}", f"={OV}!S{o}", B.F_BOLD, fmt=PCT)
        put(ws, f"G{r}", f"={OV}!U{o}", fmt=YEN)
        put(ws, f"H{r}", f"={OV}!AA{o}", fmt="#,##0")
        put(ws, f"I{r}", f"={OV}!AB{o}", B.F_BOLD, fmt=YEN)
        put(ws, f"J{r}", c["rank1"], align=B.WRAP_C)
        cell = put(ws, f"K{r}", c["sheet"], B.F_LINK)
        B.link_sheet(cell, f"'{c['sheet']}'!A1")
        ws.row_dimensions[r].height = 30
    r2 = r1 + 2 + len(D.FINALISTS) + 1
    put(ws, f"A{r2}", "年商5億の到達見込み（04シートから自動参照）", B.F_SUB, border=False)
    for j, h in enumerate(["", "1年目", "2年目", "3年目"], start=1):
        put(ws, f"{B.get_column_letter(j)}{r2 + 1}", h, B.F_HEAD, fill=B.FILL_HEAD, align=B.WRAP_C)
    u, RM = rm["uplift"], B.RM
    put(ws, f"A{r2 + 2}", "総年商（Amazon＋楽天・Yahoo!）", B.F_BOLD)
    put(ws, f"A{r2 + 3}", "総年利（営業利益・固定費前）", B.F_BOLD)
    for col, s_, p_ in (("B", "K", "N"), ("C", "L", "O"), ("D", "M", "P")):
        put(ws, f"{col}{r2 + 2}", f"={RM}!{s_}{u + 1}", B.F_BOLD, fmt=YEN)
        put(ws, f"{col}{r2 + 3}", f"={RM}!{p_}{u + 2}", B.F_BOLD, fmt=YEN)
    put(ws, f"A{r2 + 4}", "3年目の目標達成率（年商5億円比）", B.F_BOLD)
    put(ws, f"D{r2 + 4}", f"={RM}!K{u + 5}", B.F_RED, fmt=PCT)
    r3 = r2 + 6
    put(ws, f"A{r3}", "次にやること（優先順）", B.F_SUB, border=False)
    todo = ["① 天瞳の国内送料表を「前提条件」E34:E46に入力（全シート再計算）",
            "② HERO3品の1688見積（各3社）とサンプル発注。特にベビーサークル・猫ケージは梱包寸法の実測が採用条件",
            "③ Keepaで上位品の30/60/90日ランキングを確認し、各R0xシートの緑セルを埋める",
            "④ ステッパーはシリンダー耐久試験、ベビーサークルはSG基準相当の試験を工場に依頼",
            "⑤ 商標出願・JAN・PL保険"]
    for i, t in enumerate(todo):
        r = r3 + 1 + i
        ws.merge_cells(f"A{r}:K{r}")
        put(ws, f"A{r}", t, B.F_BASE, align=B.WRAP_TL, border=False)


def main(template, out):
    wb = load_workbook(template)
    tpl = wb["リサーチ表"]
    rows = D.candidates()
    models = {c["id"]: B.model(c) for c in rows}
    cands = {c["id"]: c for c in rows}
    ws_sum = wb.create_sheet("00_結論")
    ws_mk = wb.create_sheet("01_市場スキャン121")
    ws_ov = wb.create_sheet("02_商品化判断")
    ws_st = wb.create_sheet("03_1位獲得戦略")
    ws_rm = wb.create_sheet("04_年商5億ロードマップ")
    ws_as = wb.create_sheet("前提条件")
    ws_pk = wb.create_sheet("梱包明細")
    ws_ra = wb.create_sheet("生データ_Amazon")
    ws_r1 = wb.create_sheet("生データ_1688")
    ws_src = wb.create_sheet("出典")

    B.OV = "'02_商品化判断'"
    B.RM = "'04_年商5億ロードマップ'"
    B.build_assumptions(ws_as)
    rowmap, last_ov = B.build_overview(ws_ov, rows, models)
    ws_ov.title = "02_商品化判断"
    ws_ov["A1"].value = f"商品化判断（採用7品＋実データで検証して外した{len(rows) - len(D.FINALISTS)}品）"
    order = sorted(rows, key=lambda c: rowmap[c["id"]])
    pk_first = B.build_parcels(ws_pk, order)
    B.build_strategy(ws_st)
    ws_st.title = "03_1位獲得戦略"
    ws_st["A1"].value = "ランキング1位を取るための戦略（採用7品・Amazon実データ）"
    rm = B.build_roadmap(ws_rm, rowmap)
    ws_rm.title = "04_年商5億ロードマップ"
    notes = [
        "・年商5億円は1商品では届かない。HERO3品（ステッパー・ベビーサークル・宅配ボックス）を各市場の1位ライン（月600〜1,000個）近くまで伸ばし、"
        "主力2品・条件付き1品・サポート1品と楽天・Yahoo!展開で3年目に到達させる計画。",
        "・月販目標は各市場の1位の実績（購入バッジ）以下に設定。1位ラインに届かない商品は月販を下げ、色・サイズ違いで補う。",
        "・年収5,000万円（役員報酬）の原資は総年利から固定費と法人税等を引いた残り。",
        "・1年目は広告10%・導入価格期間があるため利益が薄い（1個利益（広告10%）で試算）。",
    ]
    for i, tx in enumerate(notes):
        ws_rm[f"A{rm['uplift'] + 9 + i}"].value = tx
    markets = build_market(ws_mk)
    build_summary(ws_sum, rowmap, last_ov, rm, markets)
    build_raw_amazon(ws_ra)
    build_raw_1688(ws_r1)
    B.build_sources(ws_src)
    ws_src["A1"].value = "出典（2026-10-20 Amazon.co.jp・1688を直接取得）"
    sheets = []
    for spec in D.SHEETS:
        cand = cands[spec["cand"]]
        sheets.append(B.build_product_sheet(wb, tpl, spec, cand, rowmap[cand["id"]], pk_first[cand["id"]]))
    for ws in sheets:
        ws["E9"].value = ("※ランキング30/60/90日平均はKeepa等で取得してください（緑のまま）。販売価格・購入数・BSR・評価は2026/10/20のAmazon実データ。"
                          "梱包寸法と1688原価は推定（青字）。")
    tpl.title = "テンプレート(原本)"
    wb._sheets = [ws_sum, ws_mk, ws_ov, ws_st, ws_rm] + sheets + [ws_as, ws_pk, ws_ra, ws_r1, ws_src, tpl]
    wb.active = 0
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = ws is ws_sum
    wb.save(out)
    for c in order[:12]:
        m = models[c["id"]]
        print(f"{c['id']} {m['judge']:4s} tmpl {m['tm']*100:5.1f}% real5 {m['r5']*100:5.1f}% p5 {m['p5']:6.0f} {c['name'][:30]}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
