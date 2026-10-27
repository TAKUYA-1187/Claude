# -*- coding: utf-8 -*-
"""Rebuild the 天瞳 research-table template (sheet リサーチ表) when the uploaded original is
not available: same cell positions, labels and formulas as the original; styling simplified.

usage: python3 make_template.py out.xlsx
"""
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

F = Font(name="游ゴシック", size=10)
FB = Font(name="游ゴシック", size=10, bold=True)
GREEN = PatternFill("solid", fgColor="C6EFCE")
HEAD = PatternFill("solid", fgColor="D9E1F2")
THIN = Side(style="thin", color="999999")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CEN = Alignment(horizontal="center", vertical="center", wrap_text=True)


def block(ws, top):
    """Rival block: rows top..top+6 (original rows 2-8 and 10-16)."""
    r = top
    for ref, v in ((f"E{r}", "１他社商品名"), (f"E{r + 1}", "サイト"), (f"E{r + 2}", "メーカー名"),
                   (f"H{r + 2}", "包装の寸法(cm)"), (f"L{r + 2}", "梱包重量"), (f"N{r + 2}", "CBM（体積）"),
                   (f"O{r + 2}", "50KGまで発送可能（重量）"), (f"P{r + 2}", "製品三辺合計"), (f"E{r + 3}", "ライバル販売価格"),
                   (f"H{r + 3}", "製品サイズ（㎝）"), (f"L{r + 3}", "商品重量"), (f"M{r + 3}", "50KGまで発送可能"),
                   (f"E{r + 4}", "ライバルASIN"), (f"E{r + 6}", "売れ筋ランキング - 30 days avg."),
                   (f"I{r + 6}", "売れ筋ランキング - 60 days avg."), (f"N{r + 6}", "売れ筋ランキング - 90 days avg.")):
        ws[ref] = v
        ws[ref].font = FB
    ws[f"N{r + 3}"] = f"=(I{r + 2}/100)*(J{r + 2}/100)*(K{r + 2}/100)"
    ws[f"P{r + 3}"] = f"=I{r + 2}+J{r + 2}+K{r + 2}"
    h = r + 4
    for col, v in zip("GHIJKLMNOPQRSTUVWX", ("販売価額", "仕入れ原価（元）", "レート", "円の商品原価", "国内送料", "Amazon手数料",
                                             "消費税", "関税", "１CBM", "国際運送", "粗利", "利益率", "月間売上", "販売目標個数",
                                             "月間利益", "３ヶ月平均rank", "仕入れ個数", "仕入れ個数×原価")):
        c = ws[f"{col}{h}"]
        c.value, c.font, c.fill, c.alignment, c.border = v, FB, HEAD, CEN, BOX
    v = h + 1
    vals = {"I": 23.5, "J": f"=I{v}*H{v}", "L": f"=G{v}/(100/16.5)", "M": f"=J{v}*0.1", "N": f"=J{v}*0.1", "O": 17325,
            "P": f"=O{v}*N{r + 3}", "Q": f"=G{v}-J{v}-K{v}-L{v}-M{v}-N{v}-P{v}", "R": f"=Q{v}/G{v}", "S": f"=T{v}*G{v}",
            "T": 100, "U": f"=T{v}*Q{v}", "V": 0, "W": 100, "X": f"=W{v}*P{v}+W{v}*J{v}"}
    for col in "GHIJKLMNOPQRSTUVWX":
        c = ws[f"{col}{v}"]
        c.value = vals.get(col)
        c.border = BOX
        c.alignment = CEN
    for ref in [f"F{r}", f"F{r + 1}", f"F{r + 2}", f"F{r + 3}", f"I{r + 2}", f"J{r + 2}", f"K{r + 2}", f"M{r + 2}",
                f"I{r + 3}", f"J{r + 3}", f"K{r + 3}", f"M{r + 3}", f"E{v}", f"G{v}", f"H{v}", f"K{v}", f"T{v}", f"V{v}",
                f"W{v}", f"G{r + 6}", f"K{r + 6}", f"P{r + 6}"]:
        ws[ref].fill = GREEN
    ws[f"R{v}"].number_format = "0.0%"


def main(out):
    wb = Workbook()
    ws = wb.active
    ws.title = "リサーチ表"
    ws["A1"] = "リサーチ"
    ws["A1"].font = Font(name="游ゴシック", size=20, bold=True)
    block(ws, 2)
    block(ws, 10)
    ws["F19"] = "確認事項"
    ws["F19"].font = FB
    notes = [
        ("記入項目", "緑になっている部分は入力しなければなりません。"),
        ("レート", "時期により、レートが変動されます。弊社は一定時期の平均値を取って決めております。"),
        ("消費税", "ここの消費税は仕入原価の消費税です。免税対象ではなければ、仕入原価の10％の消費税を追加して、計算してください"),
        ("梱包サイズ", "製品サイズと梱包サイズは要注意です。梱包サイズはライバルページになければ、1688で似た商品を参考しましょう。"
                     "希望原価を出す段階では梱包サイズを確認してください。"),
        ("ランキング", "ランキングの30日６０日９０日の調査は有料版ツールが必要です。初期段階はライバルのランキングを10日間くらい追跡して、"
                    "ランキングが大きく変動しないか確認してください。"),
        ("販売価格", "初期販売するとき、ライバルよりやすい価格でスタートします。売れてから徐々に値上げします"),
        ("ライバル販売価格", "ライバルの販売価格はクーポンがある分でしたら、割引を引いた価格になります。"),
        ("国際送料", "国際送料の単価は発注合計の体積を計算して、2㎥に超えば単価は17325円になります。2㎥に超えなければ単価は33000円になります。"),
        ("特大サイズ", "梱包サイズは三辺合計260cm超える場合は2個口か3個口に分けて発送できるようにします"),
    ]
    for i, (g, h) in enumerate(notes):
        r = 20 + i
        ws[f"F{r}"], ws[f"G{r}"], ws[f"H{r}"] = i + 1, g, h
    qs = [(30, "１：競合セラーのレビューを確認した上で、商品のメリットをまとめ\n\n",
           "例えば、折りたたみ可能、軽量（5kg）、耐荷重100kg ，組立不要，などのポイントが挙げられます。"),
          (35, "２：競合セラーのレビューを確認した上で、商品の悪いのところをまとめ\n\n",
           "例えば、説明書がわかりにくく、不親切，カスタマーサービスに対する不満、品質にばらつきがある・壊れやすいなど"),
          (40, "3：自社サンプルにはどのようなメリットがあるのか、\nまた、上記の低評価レビューに挙げられている問題点を解決できるかどうか",
           "例えば、詳細な取扱説明書，組み立て動画あり、また、製品構造を強化して耐荷重を上げたり、より品質の高いサンプルを選定すること\n"),
          (46, "重要４：レビューの中から、お客様の潜在的なニーズを抽出することは可能？",
           "例えば、「もっと大きいサイズがあればよかった」「この色だったら好きなのに」「もう少し軽ければ嬉しい」「組み立てがもっと簡単ならよかった」など")]
    for r, q, ex in qs:
        ws[f"A{r}"], ws[f"B{r}"] = q, ex
        ws[f"A{r}"].font = FB
        ws[f"A{r}"].alignment = ws[f"B{r}"].alignment = WRAP
        for k in range(1, 5):
            ws[f"B{r + k}"].fill = GREEN
    ws["A53"] = "フィードバック"
    ws["A53"].font = FB
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 50
    for col in "EFGHIJKLMNOPQRSTUVWX":
        ws.column_dimensions[col].width = 12
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None and c.font == Font():
                c.font = F
    wb.save(out)


if __name__ == "__main__":
    main(sys.argv[1])
